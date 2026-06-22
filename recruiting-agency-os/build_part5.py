"""Recruiting Agency OS — Part 5: PDFs + PPTX Canva Templates + manifest + ZIP"""
import os, csv, json, zipfile
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.units import cm
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.text import PP_ALIGN

BASE = "/home/user/oqul-phase55-production/recruiting-agency-os/Ultimate_Recruiting_Agency_Operating_System/"

NAV_RGB = (15, 52, 96)
ACC_RGB = (233, 69, 96)
GLD_RGB = (245, 166, 35)
GRN_RGB = (39, 174, 96)
WHT_RGB = (255, 255, 255)
LGR_RGB = (248, 249, 250)

PNAV = PRGB(15, 52, 96)
PACC = PRGB(233, 69, 96)
PGLD = PRGB(245, 166, 35)
PGRN = PRGB(39, 174, 96)
PWHT = PRGB(255, 255, 255)

# ── PDFs ──────────────────────────────────────────────────────────────────────
def make_pdf(path, title, sections):
    doc = SimpleDocTemplate(path, pagesize=A4,
        leftMargin=2*cm, rightMargin=2*cm, topMargin=2*cm, bottomMargin=2*cm)
    styles = getSampleStyleSheet()
    NAV_C = colors.Color(15/255, 52/255, 96/255)
    ACC_C = colors.Color(233/255, 69/255, 96/255)

    title_style = ParagraphStyle('T', parent=styles['Normal'],
        fontSize=20, textColor=NAV_C, spaceAfter=6, fontName='Helvetica-Bold')
    h1_style = ParagraphStyle('H1', parent=styles['Normal'],
        fontSize=13, textColor=NAV_C, spaceAfter=4, spaceBefore=14, fontName='Helvetica-Bold')
    body_style = ParagraphStyle('B', parent=styles['Normal'],
        fontSize=10, spaceAfter=4, leading=14)
    bullet_style = ParagraphStyle('BU', parent=styles['Normal'],
        fontSize=10, spaceAfter=3, leading=14, leftIndent=15, bulletIndent=5)

    story = [Paragraph(title, title_style),
             HRFlowable(width="100%", thickness=2, color=ACC_C, spaceAfter=10)]
    for heading, items in sections:
        story.append(Paragraph(heading, h1_style))
        for item in items:
            if item.startswith('• '):
                story.append(Paragraph(item, bullet_style))
            else:
                story.append(Paragraph(item, body_style))
    doc.build(story)
    print(f"  ✓ {os.path.relpath(path, BASE)}")

make_pdf(BASE + "00_START_HERE/START_HERE_Implementation_Guide.pdf",
    "Ultimate Recruiting Agency Operating System — Implementation Guide",
    [
        ("WELCOME", [
            "Congratulations on your investment in the Ultimate Recruiting Agency Operating System. This guide walks you through implementation in a logical, step-by-step sequence so you can use every resource effectively.",
        ]),
        ("WHAT'S INCLUDED", [
            "• 01_BUSINESS_DEVELOPMENT — Cold outreach, proposals, follow-up sequences, pricing guide, BD CRM",
            "• 02_CLIENT_ONBOARDING — Welcome emails, job order intake, kickoff scripts, expectations guide",
            "• 03_CANDIDATE_SOURCING — LinkedIn playbook, boolean search library, job board strategy",
            "• 04_CANDIDATE_MANAGEMENT — Screening scripts, assessment framework, communication templates",
            "• 05_CLIENT_MANAGEMENT — Candidate submittal, weekly updates, feedback collection, difficult conversations",
            "• 06_COMPLIANCE_LEGAL — Service agreement, anti-discrimination policy, EEOC checklist, record retention",
            "• 07_PLACEMENT_PROCESS — Offer management, salary negotiation, handoff, post-placement follow-up",
            "• 08_BUSINESS_OPERATIONS — Goals workbook, team management, recruiter training manual, KPI dashboard",
            "• 09_MARKETING_GROWTH — Content strategy, LinkedIn playbook, referral system, personal brand guide",
            "• 10_NOTION_WORKSPACE — 8 databases, setup guide, import-ready CSV files",
            "• 11_BONUSES — 365 captions, 100 content ideas, success framework, growth roadmap",
            "• 12_CANVA_IMPORTABLE_TEMPLATES — 4 PPTX files for Canva (proposal, candidate profile, job post, social pack)",
        ]),
        ("IMPLEMENTATION SEQUENCE", [
            "Week 1 — Foundation: Review 06_COMPLIANCE_LEGAL. Customize and sign contracts. Set up Notion workspace using 10_NOTION_WORKSPACE files.",
            "Week 2 — Business Development: Customize 01_BUSINESS_DEVELOPMENT templates. Launch cold outreach campaign. Set up Business_Development_CRM.xlsx.",
            "Week 3 — Sourcing Setup: Customize boolean searches in 03_CANDIDATE_SOURCING. Launch LinkedIn strategy. Create job board accounts.",
            "Week 4 — Full Operations: Begin using 04_CANDIDATE_MANAGEMENT screening scripts. Use 02_CLIENT_ONBOARDING for first client. Track all activity in XLSX dashboards.",
            "Ongoing — Growth: Post 3x/week on LinkedIn using 09_MARKETING_GROWTH and 11_BONUSES captions. Track KPIs weekly in Business_KPI_Dashboard.xlsx.",
        ]),
        ("CANVA TEMPLATE IMPORT", [
            "To edit PPTX files in Canva:",
            "• Go to canva.com → Click 'Create a design'",
            "• Select 'Import file' and upload the .pptx file",
            "• Canva converts it to an editable design",
            "• Edit colors, fonts, and text to match your brand",
        ]),
        ("NOTION SETUP QUICK START", [
            "• Open 10_NOTION_WORKSPACE/NOTION_WORKSPACE_SETUP.md for full instructions",
            "• Import each CSV file into Notion as a new database",
            "• Create relations between Candidates, Job Orders, Clients, and Placements",
            "• Build your Dashboard page using linked database views",
        ]),
        ("PRICING REFERENCE", [
            "• Original Price: €149",
            "• Launch Price: €39",
            "• Share your results with us — we love seeing agencies grow with these systems!",
        ]),
    ])

make_pdf(BASE + "06_COMPLIANCE_LEGAL/Legal_Template_Use_Guide.pdf",
    "Legal Template Use Guide — Recruiting Agency OS",
    [
        ("IMPORTANT NOTICE", [
            "The legal documents in this operating system are generic educational templates only. They are not legal advice. Always have a qualified attorney review any contract before commercial use.",
        ]),
        ("DOCUMENTS INCLUDED", [
            "• Recruiting_Service_Agreement.docx — Core client contract covering fees, guarantees, confidentiality",
            "• Anti_Discrimination_Policy.docx — Internal policy for EEOC compliance",
            "• EEOC_Compliance_Checklist.docx — Ongoing compliance reference",
            "• Record_Retention_Policy.docx — Data retention schedule and requirements",
        ]),
        ("CUSTOMIZATION STEPS", [
            "1. Open each document in Microsoft Word or Google Docs",
            "2. Replace all [bracketed placeholders] with your specific information",
            "3. Review all percentage figures, timelines, and terms",
            "4. Have a licensed attorney in your jurisdiction review before use",
            "5. Keep signed copies of all agreements in a secure location",
        ]),
        ("KEY LEGAL AREAS FOR RECRUITING AGENCIES", [
            "• Fee agreements: Contingency vs. retained; payment timing; guarantee terms",
            "• EEOC/Anti-discrimination: Prohibited questions; disparate impact; record keeping",
            "• Privacy: FCRA for background checks; GDPR/CCPA for candidate data",
            "• Confidentiality: Candidate data protection; client confidential information",
            "• Intellectual property: Ownership of candidate databases and sourced leads",
        ]),
        ("WHEN TO CONSULT AN ATTORNEY", [
            "• Before your first client signs a contract",
            "• When entering a new state with different employment laws",
            "• When a dispute arises with a client or candidate",
            "• When handling background checks or sensitive candidate data",
            "• When adding employees or independent contractors to your team",
        ]),
    ])

make_pdf(BASE + "10_NOTION_WORKSPACE/Notion_Workspace_Setup_Guide.pdf",
    "Notion Workspace Setup Guide — Recruiting Agency OS",
    [
        ("OVERVIEW", [
            "Your Notion workspace connects all aspects of your recruiting agency — from BD pipeline to candidate management and placements. This guide walks you through setting it up in under 2 hours.",
        ]),
        ("DATABASE OVERVIEW", [
            "8 databases are included in your CSV import files:",
        ]),
        ("DATABASE STRUCTURE", [
            "• Candidates_Database.csv — Full candidate tracking with stage, score, and client match",
            "• Clients_Database.csv — Client accounts, active roles, and placement history",
            "• Job_Orders_Database.csv — Active searches with priority, status, and timeline",
            "• Placements_Database.csv — Completed placements with guarantee tracking",
            "• Tasks_Database.csv — Daily task management linked to searches and clients",
            "• Business_Development_Database.csv — BD pipeline with stage and follow-up tracking",
            "• Content_Calendar_Database.csv — LinkedIn and marketing content planning",
            "• SOP_Library.csv — All agency SOPs with version control",
        ]),
        ("IMPORT STEPS", [
            "1. Open Notion and create a new page titled 'Recruiting Agency OS'",
            "2. Click the '+' button → 'Import' → 'CSV'",
            "3. Upload Candidates_Database.csv — rename to 'Candidates'",
            "4. Repeat for all 8 CSV files",
            "5. Open NOTION_WORKSPACE_SETUP.md for advanced setup (relations, formulas, views)",
        ]),
        ("RECOMMENDED VIEWS TO CREATE", [
            "• Candidates: Board view grouped by Stage",
            "• Job Orders: Board view grouped by Status; Calendar view by Start Date",
            "• Tasks: Board view grouped by Priority; Calendar view by Due Date",
            "• BD Pipeline: Board view grouped by Stage",
            "• Content Calendar: Calendar view by Publish Date",
        ]),
        ("WEEKLY WORKFLOW IN NOTION", [
            "Monday: Review Tasks for the week; check BD pipeline for follow-ups due",
            "Daily: Log all candidate and client communications as tasks",
            "Friday: Update all Job Order statuses; send weekly client updates",
            "Monthly: Review Placements database for guarantee expiration alerts",
        ]),
    ])

print("✓ PDFs done")

# ── PPTX CANVA TEMPLATES ──────────────────────────────────────────────────────
def new_prs():
    prs = Presentation()
    prs.slide_width  = Inches(13.33)
    prs.slide_height = Inches(7.5)
    return prs

def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])

def rect(slide, l, t, w, h, r, g, b):
    from pptx.util import Inches, Emu
    shape = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = PRGB(r, g, b)
    shape.line.fill.background()
    return shape

def txt(slide, text, l, t, w, h, sz=18, bold=False, color=PWHT, align=PP_ALIGN.LEFT, italic=False):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = text
    run.font.size = Pt(sz); run.font.bold = bold
    run.font.color.rgb = color; run.font.italic = italic

def accent_line(slide, l, t, w):
    ln = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(0.04))
    ln.fill.solid(); ln.fill.fore_color.rgb = PACC
    ln.line.fill.background()

# 1. Client Proposal
prs = new_prs()

# Slide 1 — Cover
sl = blank_slide(prs)
rect(sl, 0, 0, 13.33, 7.5, *NAV_RGB)
rect(sl, 0, 0, 4.5, 7.5, *ACC_RGB)
txt(sl, "RECRUITING PROPOSAL", 5, 1.5, 8, 0.8, 36, True, PWHT, PP_ALIGN.LEFT)
txt(sl, "[Client Company Name]", 5, 2.5, 8, 0.6, 22, False, PWHT, PP_ALIGN.LEFT)
txt(sl, "Prepared by [Agency Name]  |  [Date]", 5, 3.2, 8, 0.5, 14, False, PRGB(180,180,180), PP_ALIGN.LEFT)
txt(sl, "[Role Title]", 0.3, 5.5, 3.9, 0.8, 26, True, PWHT, PP_ALIGN.LEFT)
txt(sl, "Search", 0.3, 6.3, 3.9, 0.5, 14, False, PRGB(200,200,200), PP_ALIGN.LEFT)

# Slide 2 — Understanding
sl = blank_slide(prs)
rect(sl, 0, 0, 13.33, 7.5, *LGR_RGB)
rect(sl, 0, 0, 13.33, 1.2, *NAV_RGB)
txt(sl, "OUR UNDERSTANDING OF YOUR NEEDS", 0.4, 0.25, 12, 0.7, 22, True, PWHT, PP_ALIGN.LEFT)
items = [
    ("Role:", "[Job Title] — [Department]"),
    ("Salary Range:", "$[Min] – $[Max]"),
    ("Location:", "[City] | [Remote/Hybrid/On-Site]"),
    ("Start Date:", "[Target Date]"),
    ("Must-Haves:", "[Top 3 requirements from intake]"),
    ("Culture:", "[Key culture notes]"),
]
for i, (label, val) in enumerate(items):
    y = 1.5 + i * 0.85
    rect(sl, 0.4, y, 12.5, 0.7, *WHT_RGB)
    txt(sl, label, 0.6, y + 0.1, 2.5, 0.5, 12, True, PNAV)
    txt(sl, val, 3.2, y + 0.1, 9.5, 0.5, 12, False, PRGB(50,50,50))

# Slide 3 — Our Approach
sl = blank_slide(prs)
rect(sl, 0, 0, 13.33, 7.5, *NAV_RGB)
txt(sl, "OUR APPROACH", 0.5, 0.4, 12, 0.7, 26, True, PWHT, PP_ALIGN.LEFT)
accent_line(sl, 0.5, 1.2, 4)
steps = [
    ("01", "Week 1", "Deep search of 20–40 potential candidates across LinkedIn, database, and niche sources."),
    ("02", "Week 1–2", "Phone screening against your criteria. Competency interviews and culture assessment."),
    ("03", "Week 2", "Deliver 3–5 fully profiled candidates with recruiter assessment and notes."),
    ("04", "Ongoing", "Interview scheduling, candidate prep, offer management, and onboarding handoff."),
]
for i, (num, time, desc) in enumerate(steps):
    x = 0.4 + i * 3.2
    rect(sl, x, 1.5, 2.9, 4.8, *ACC_RGB)
    txt(sl, num, x+0.2, 1.7, 2.5, 0.7, 32, True, PWHT)
    txt(sl, time, x+0.2, 2.5, 2.5, 0.4, 13, True, PWHT)
    txt(sl, desc, x+0.2, 3.1, 2.5, 2.5, 11, False, PWHT)

# Slide 4 — Fee Structure
sl = blank_slide(prs)
rect(sl, 0, 0, 13.33, 7.5, *LGR_RGB)
rect(sl, 0, 0, 13.33, 1.2, *NAV_RGB)
txt(sl, "FEE STRUCTURE & GUARANTEE", 0.4, 0.25, 12, 0.7, 22, True, PWHT, PP_ALIGN.LEFT)
fees = [
    ("FEE TYPE", "[X]% of First-Year Base Salary"),
    ("PAYMENT", "Due within 30 days of candidate start date"),
    ("GUARANTEE", "[90]-day replacement at no additional fee"),
    ("SCOPE", "One active search per Job Order"),
    ("EXCLUSIVITY", "[ ] Exclusive  [ ] Contingency  — confirmed at kickoff"),
]
for i, (label, val) in enumerate(fees):
    y = 1.5 + i * 1.0
    rect(sl, 0.4, y, 4, 0.8, *NAV_RGB)
    rect(sl, 4.5, y, 8.4, 0.8, *WHT_RGB)
    txt(sl, label, 0.5, y+0.15, 3.8, 0.5, 12, True, PWHT)
    txt(sl, val, 4.7, y+0.15, 8, 0.5, 12, False, PRGB(50,50,50))

# Slide 5 — Why Us
sl = blank_slide(prs)
rect(sl, 0, 0, 13.33, 7.5, *NAV_RGB)
txt(sl, "WHY [AGENCY NAME]", 0.5, 0.4, 12, 0.7, 26, True, PWHT, PP_ALIGN.LEFT)
accent_line(sl, 0.5, 1.2, 4)
points = [
    ("[X] Years", "Specializing in [niche] recruiting"),
    ("[X] Days", "Average time-to-fill"),
    ("[X]%", "Placement retention at 12 months"),
    ("[X]+", "Pre-vetted candidates in database"),
]
for i, (stat, label) in enumerate(points):
    x = 0.4 + i * 3.2
    rect(sl, x, 1.6, 2.9, 3.5, *PRGB(255,255,255).__class__(*WHT_RGB))
    txt(sl, stat, x+0.2, 1.9, 2.5, 1, 30, True, PNAV)
    txt(sl, label, x+0.2, 3.0, 2.5, 1, 13, False, PRGB(70,70,70))
txt(sl, '"We specialize exclusively in [niche]. Every candidate is personally screened by our team."', 0.5, 5.4, 12.3, 1, 15, False, PRGB(180,200,220))

# Slide 6 — Next Steps
sl = blank_slide(prs)
rect(sl, 0, 0, 13.33, 7.5, *LGR_RGB)
rect(sl, 0, 0, 13.33, 1.2, *NAV_RGB)
txt(sl, "NEXT STEPS", 0.4, 0.25, 12, 0.7, 22, True, PWHT, PP_ALIGN.LEFT)
steps6 = [
    ("01", "Review & Sign Service Agreement", "Included in this proposal package."),
    ("02", "Complete Job Order Intake Form", "15-minute form to capture full role requirements."),
    ("03", "Kickoff Call (30 min)", "Align on search criteria, process, and communication cadence."),
    ("04", "First Candidate Profiles", "Delivered within [5–7] business days of kickoff."),
]
for i, (num, title, desc) in enumerate(steps6):
    y = 1.5 + i * 1.35
    rect(sl, 0.4, y, 0.8, 1.1, *NAV_RGB)
    txt(sl, num, 0.45, y+0.2, 0.7, 0.6, 20, True, PWHT, PP_ALIGN.CENTER)
    txt(sl, title, 1.4, y+0.1, 10, 0.4, 14, True, PNAV)
    txt(sl, desc, 1.4, y+0.55, 10, 0.4, 12, False, PRGB(80,80,80))

prs.save(BASE + "12_CANVA_IMPORTABLE_TEMPLATES/Canva_Import_Client_Proposal.pptx")
print("  ✓ 12_CANVA_IMPORTABLE_TEMPLATES/Canva_Import_Client_Proposal.pptx")

# 2. Candidate Profile
prs = new_prs()

sl = blank_slide(prs)
rect(sl, 0, 0, 13.33, 7.5, *NAV_RGB)
rect(sl, 9.5, 0, 3.83, 7.5, *ACC_RGB)
txt(sl, "CANDIDATE", 0.5, 0.8, 8, 0.6, 14, False, PRGB(180,200,220))
txt(sl, "PROFILE", 0.5, 1.4, 8, 1, 52, True, PWHT)
txt(sl, "[Candidate Full Name]", 0.5, 2.6, 8, 0.7, 24, False, PWHT)
txt(sl, "[Current Role] at [Current Company]", 0.5, 3.3, 8, 0.5, 16, False, PRGB(180,200,220))
txt(sl, "[Agency Name]  |  Confidential", 9.7, 6.5, 3.3, 0.5, 11, False, PWHT)

sl = blank_slide(prs)
rect(sl, 0, 0, 13.33, 7.5, *LGR_RGB)
rect(sl, 0, 0, 13.33, 1.2, *NAV_RGB)
txt(sl, "CANDIDATE OVERVIEW", 0.4, 0.25, 12, 0.7, 22, True, PWHT, PP_ALIGN.LEFT)
fields = [
    ("Current Title", "[Title] at [Company] — [X] years"),
    ("Location", "[City, State]  |  [Remote/Relocation preference]"),
    ("Education", "[Degree], [Institution]"),
    ("Availability", "[Date]  |  Notice: [X weeks]"),
    ("Compensation", "Current: $[X]  |  Seeking: $[X] – $[X]"),
]
for i, (label, val) in enumerate(fields):
    y = 1.4 + i * 1.0
    rect(sl, 0.4, y, 3.5, 0.8, *NAV_RGB)
    rect(sl, 4.0, y, 9, 0.8, *WHT_RGB)
    txt(sl, label, 0.5, y+0.15, 3.3, 0.5, 12, True, PWHT)
    txt(sl, val, 4.15, y+0.15, 8.7, 0.5, 12, False, PRGB(50,50,50))

sl = blank_slide(prs)
rect(sl, 0, 0, 13.33, 7.5, *NAV_RGB)
rect(sl, 0, 0, 13.33, 1.2, *ACC_RGB)
txt(sl, "EXPERIENCE SUMMARY & KEY QUALIFICATIONS", 0.4, 0.25, 12, 0.7, 18, True, PWHT, PP_ALIGN.LEFT)
txt(sl, "EXPERIENCE SUMMARY", 0.4, 1.4, 12, 0.5, 14, True, PGLD)
txt(sl, "[3–4 sentences highlighting most relevant experience for this specific role. Focus on measurable achievements and direct relevance to the client's requirements.]", 0.4, 2.0, 12.5, 1.5, 13, False, PWHT)
txt(sl, "KEY QUALIFICATIONS", 0.4, 3.8, 12, 0.5, 14, True, PGLD)
quals = ["[Requirement 1]: [Evidence from candidate background]",
         "[Requirement 2]: [Evidence from candidate background]",
         "[Requirement 3]: [Evidence from candidate background]"]
for i, q in enumerate(quals):
    txt(sl, f"• {q}", 0.4, 4.4 + i*0.7, 12.5, 0.6, 12, False, PRGB(200,220,240))

sl = blank_slide(prs)
rect(sl, 0, 0, 13.33, 7.5, *LGR_RGB)
rect(sl, 0, 0, 13.33, 1.2, *NAV_RGB)
txt(sl, "RECRUITER ASSESSMENT & INTEREST LEVEL", 0.4, 0.25, 12, 0.7, 20, True, PWHT, PP_ALIGN.LEFT)
rect(sl, 0.4, 1.4, 8, 4.5, *WHT_RGB)
txt(sl, "RECRUITER'S ASSESSMENT", 0.6, 1.6, 7.5, 0.4, 13, True, PNAV)
txt(sl, "[1–2 sentences summarizing overall fit, standout strengths, and any minor considerations. Be specific and direct about why this candidate is a strong match.]", 0.6, 2.2, 7.5, 2.0, 12, False, PRGB(50,50,50))
txt(sl, "RECOMMENDATION:", 0.6, 4.5, 7.5, 0.4, 13, True, PACC)
txt(sl, "[ ] Highly Recommended   [ ] Recommended   [ ] With Reservations", 0.6, 5.0, 7.5, 0.4, 12, False, PRGB(50,50,50))
rect(sl, 9, 1.4, 3.9, 4.5, *NAV_RGB)
txt(sl, "INTEREST\nLEVEL", 9.2, 1.8, 3.5, 1, 16, True, PWHT, PP_ALIGN.CENTER)
txt(sl, "8/10", 9.2, 3.0, 3.5, 1.2, 52, True, PGLD, PP_ALIGN.CENTER)
txt(sl, "[Brief reason for score]", 9.2, 4.4, 3.5, 1, 12, False, PRGB(180,200,220), PP_ALIGN.CENTER)

prs.save(BASE + "12_CANVA_IMPORTABLE_TEMPLATES/Canva_Import_Candidate_Profile.pptx")
print("  ✓ 12_CANVA_IMPORTABLE_TEMPLATES/Canva_Import_Candidate_Profile.pptx")

# 3. Job Opportunity (candidate-facing)
prs = new_prs()

sl = blank_slide(prs)
rect(sl, 0, 0, 13.33, 7.5, *NAV_RGB)
rect(sl, 0, 5.5, 13.33, 2, *ACC_RGB)
txt(sl, "EXCITING OPPORTUNITY", 0.8, 1.0, 12, 0.7, 14, False, PRGB(180,200,220))
txt(sl, "[Job Title]", 0.8, 1.8, 11, 1.1, 44, True, PWHT)
txt(sl, "[Company Type] | [Location] | [Compensation Range]", 0.8, 3.0, 11, 0.6, 18, False, PRGB(180,200,220))
txt(sl, "Confidential Search  •  [Agency Name]", 0.8, 5.7, 11, 0.5, 14, False, PWHT)

sl = blank_slide(prs)
rect(sl, 0, 0, 13.33, 7.5, *LGR_RGB)
rect(sl, 0, 0, 13.33, 1.2, *NAV_RGB)
txt(sl, "ROLE OVERVIEW", 0.4, 0.25, 12, 0.7, 22, True, PWHT, PP_ALIGN.LEFT)
details = [
    ("Role", "[Job Title]"),
    ("Company", "[Company Description — 1 sentence]"),
    ("Location", "[City / Remote / Hybrid]"),
    ("Salary", "$[Min] – $[Max] base + [bonus/equity]"),
    ("Benefits", "[Key benefits: healthcare, PTO, 401k, etc.]"),
    ("Start Date", "[Target Date]"),
]
for i, (k, v) in enumerate(details):
    y = 1.4 + i * 0.9
    rect(sl, 0.4, y, 3, 0.75, *NAV_RGB)
    rect(sl, 3.5, y, 9.4, 0.75, *WHT_RGB)
    txt(sl, k, 0.55, y+0.15, 2.8, 0.45, 12, True, PWHT)
    txt(sl, v, 3.65, y+0.15, 9.1, 0.45, 12, False, PRGB(50,50,50))

sl = blank_slide(prs)
rect(sl, 0, 0, 13.33, 7.5, *NAV_RGB)
rect(sl, 0, 0, 13.33, 1.2, *ACC_RGB)
txt(sl, "WHAT YOU'LL DO & WHAT YOU'LL NEED", 0.4, 0.25, 12, 0.7, 20, True, PWHT, PP_ALIGN.LEFT)
txt(sl, "KEY RESPONSIBILITIES", 0.4, 1.4, 6, 0.4, 14, True, PGLD)
resp = ["• [Responsibility 1]","• [Responsibility 2]","• [Responsibility 3]","• [Responsibility 4]"]
for i, r in enumerate(resp):
    txt(sl, r, 0.4, 2.0 + i*0.8, 6, 0.65, 12, False, PWHT)
txt(sl, "REQUIREMENTS", 7, 1.4, 6, 0.4, 14, True, PGLD)
reqs = ["• [X]+ years of experience","• [Key skill]","• [Industry background]","• [Nice-to-have]"]
for i, r in enumerate(reqs):
    txt(sl, r, 7, 2.0 + i*0.8, 6, 0.65, 12, False, PWHT)

sl = blank_slide(prs)
rect(sl, 0, 0, 13.33, 7.5, *LGR_RGB)
rect(sl, 0, 0, 13.33, 1.2, *NAV_RGB)
txt(sl, "NEXT STEPS — INTERESTED?", 0.4, 0.25, 12, 0.7, 22, True, PWHT, PP_ALIGN.LEFT)
steps4 = [("Step 1","15-Minute Call","Tell me about your background and what you're looking for."),
           ("Step 2","Role Deep-Dive","If there's mutual interest, full role details shared."),
           ("Step 3","Client Introduction","Profile submitted confidentially with your permission."),
           ("Step 4","Interview Process","We manage scheduling, prep, and feedback throughout.")]
for i, (num, title, desc) in enumerate(steps4):
    x = 0.4 + i * 3.2
    rect(sl, x, 1.6, 2.9, 4.5, *NAV_RGB)
    txt(sl, num, x+0.2, 1.8, 2.5, 0.5, 12, True, PACC)
    txt(sl, title, x+0.2, 2.4, 2.5, 0.6, 14, True, PWHT)
    txt(sl, desc, x+0.2, 3.2, 2.5, 2, 11, False, PRGB(180,200,220))
txt(sl, "Ready to learn more? Contact: [Your Name]  |  [email]  |  [phone]", 0.4, 6.6, 12.5, 0.5, 13, False, PNAV, PP_ALIGN.CENTER)

prs.save(BASE + "12_CANVA_IMPORTABLE_TEMPLATES/Canva_Import_Job_Opportunity.pptx")
print("  ✓ 12_CANVA_IMPORTABLE_TEMPLATES/Canva_Import_Job_Opportunity.pptx")

# 4. Social Post Pack
prs = new_prs()
posts = [
    ("MYTH", "Agencies just send you random resumes.", "The best recruiters screen every candidate personally before submitting. Quality over volume, always.", NAV_RGB),
    ("HIRING TIP", "Speed is your competitive advantage.", "Top candidates are off the market in 7 days. Your fastest competitors are your biggest threat.", ACC_RGB),
    ("CAREER ADVICE", "Negotiate every offer.", "Most companies expect it. You could leave 10–20% on the table by saying nothing. Know your worth.", NAV_RGB),
    ("INDUSTRY STAT", "70% of top hires never see your job post.", "They need to be found — not waiting on Indeed. That's what great recruiters do.", ACC_RGB),
    ("QUESTION", "What's the #1 trait you look for in a new hire?", "Drop your answer below. I'll share what we hear most from top hiring managers.", NAV_RGB),
    ("CLIENT WIN", "Filled a VP Sales role in 9 days.", "Client had been searching internally for 4 months. We delivered 3 qualified profiles in 5 days. Offer accepted Day 9.", ACC_RGB),
    ("CHECKLIST", "Before you accept any job offer:", "• Research the team on LinkedIn\n• Ask about 30/60/90 day expectations\n• Confirm the comp structure in writing\n• Check Glassdoor recent reviews\n• Meet your direct manager", NAV_RGB),
    ("CTA", "Hiring in [niche] this quarter?", "We specialize in [role type] placements. Average time-to-fill: [X] days. DM me — let's talk.", ACC_RGB),
]
for label, headline, body, bg in posts:
    sl = blank_slide(prs)
    rect(sl, 0, 0, 13.33, 7.5, *bg)
    rect(sl, 0, 0, 13.33, 1.2, *ACC_RGB if bg == NAV_RGB else NAV_RGB)
    txt(sl, label, 0.4, 0.25, 12, 0.7, 14, True, PWHT, PP_ALIGN.LEFT)
    txt(sl, headline, 0.5, 1.5, 12.3, 1.5, 30, True, PWHT, PP_ALIGN.LEFT)
    txt(sl, body, 0.5, 3.4, 12.3, 3, 16, False, PRGB(200,220,240), PP_ALIGN.LEFT)
    txt(sl, "[Agency Name]  |  #recruiting #hiring", 0.5, 6.8, 12.3, 0.4, 11, False, PRGB(150,170,190))

prs.save(BASE + "12_CANVA_IMPORTABLE_TEMPLATES/Canva_Import_Social_Post_Pack.pptx")
print("  ✓ 12_CANVA_IMPORTABLE_TEMPLATES/Canva_Import_Social_Post_Pack.pptx")
print("✓ PPTX done")

# ── ASSET MANIFEST ────────────────────────────────────────────────────────────
all_files = []
for root, dirs, files in os.walk(BASE):
    for fn in sorted(files):
        rel = os.path.relpath(os.path.join(root, fn), BASE)
        folder = rel.split(os.sep)[0]
        ext = fn.rsplit('.', 1)[-1].upper()
        all_files.append((rel, folder, ext))

all_files.sort(key=lambda x: x[0])

with open(BASE + "00_START_HERE/Asset_Manifest.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(["File Path","Folder","Format","Editable In"])
    fmt_map = {"DOCX":"Microsoft Word / Google Docs","XLSX":"Microsoft Excel / Google Sheets",
               "PDF":"Adobe Acrobat (view)","CSV":"Excel / Google Sheets / Notion",
               "PPTX":"Upload to Canva / PowerPoint","MD":"Notion / Any text editor",
               "TXT":"Any text editor","JSON":"Any text editor"}
    for rel, folder, ext in all_files:
        w.writerow([rel, folder, ext, fmt_map.get(ext, ext)])

manifest_data = {
    "product": "Ultimate Recruiting Agency Operating System",
    "version": "1.0",
    "total_files": len(all_files),
    "folders": 13,
    "pricing": {"original": "€149", "launch": "€39"},
    "formats": list(set(ext for _, _, ext in all_files)),
    "files": [{"path": rel, "folder": folder, "format": ext} for rel, folder, ext in all_files]
}
with open(BASE + "00_START_HERE/Asset_Manifest.json", "w") as f:
    json.dump(manifest_data, f, indent=2)

print(f"  ✓ Manifest: {len(all_files)} files")

# ── ZIP ────────────────────────────────────────────────────────────────────────
zip_path = "/home/user/oqul-phase55-production/recruiting-agency-os/BUYER_DOWNLOAD_RecruitingAgencyOS.zip"
with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zf:
    for root, dirs, files in os.walk(BASE):
        for fn in files:
            fp = os.path.join(root, fn)
            arc = os.path.relpath(fp, os.path.dirname(BASE))
            zf.write(fp, arc)

size_mb = os.path.getsize(zip_path) / 1024 / 1024
print(f"\n✅ ZIP created: {zip_path}")
print(f"   Size: {size_mb:.1f} MB | Files: {len(all_files)}")
