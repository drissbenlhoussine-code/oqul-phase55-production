"""Bookkeeping Business OS — Part 2: folders 04–07 + XLSX trackers"""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

BASE = "/home/user/oqul-phase55-production/bookkeeping-business-os/Ultimate_Bookkeeping_Business_Operating_System/"

NAV="0F3460"; ACC="E94560"; GLD="F5A623"; GRN="27AE60"; WHT="FFFFFF"; LGR="F8F9FA"

def hf(hex_): return PatternFill("solid", fgColor=hex_)
def bf(bold=True, sz=11, color="000000"): return Font(bold=bold, size=sz, color=color)
def al(h="center", v="center"): return Alignment(horizontal=h, vertical=v, wrap_text=True)
def thin(): s=Side(style='thin',color='CCCCCC'); return Border(left=s,right=s,top=s,bottom=s)
def hrow(ws,row,cols,texts,bg=NAV,fg=WHT):
    for col,text in zip(cols,texts):
        c=ws.cell(row=row,column=col,value=text)
        c.fill=hf(bg); c.font=bf(True,11,fg); c.alignment=al(); c.border=thin()
def drow(ws,row,cols,vals,bg=WHT):
    for col,val in zip(cols,vals):
        c=ws.cell(row=row,column=col,value=val)
        c.fill=hf(bg); c.font=bf(False,10); c.alignment=al("left"); c.border=thin()

def doc(filename, title, subtitle, sections):
    d = Document()
    t = d.add_paragraph(title); t.style = d.styles['Normal']
    t.runs[0].bold = True; t.runs[0].font.size = Pt(14)
    t.runs[0].font.color.rgb = RGBColor(0x0F,0x34,0x60)
    if subtitle:
        s = d.add_paragraph(subtitle); s.style = d.styles['Normal']
        s.runs[0].font.size = Pt(10)
        s.runs[0].font.color.rgb = RGBColor(0x7F,0x8C,0x8D)
    for section in sections:
        if isinstance(section, str):
            d.add_paragraph(section); continue
        heading, items = section
        h = d.add_heading(heading, level=1)
        h.runs[0].font.color.rgb = RGBColor(0x0F,0x34,0x60)
        for item in items:
            if isinstance(item, tuple) and item[0] == '•':
                d.add_paragraph(item[1], style='List Bullet')
            else:
                d.add_paragraph(str(item))
    d.save(BASE + filename)
    print(f"  ✓ {filename}")

# ── 04_PAYROLL_SERVICES ───────────────────────────────────────────────────────
doc("04_PAYROLL_SERVICES/Payroll_Processing_SOP.docx",
    "Payroll Processing SOP",
    "Bookkeeping Business Operating System | Payroll Services",
    [
        ("PAYROLL SCHEDULE OVERVIEW", [
            "Establish clear payroll dates at onboarding. Most clients run weekly, bi-weekly, semi-monthly, or monthly payroll. Missing a payroll date creates legal liability — track deadlines in your calendar with 2-day advance reminders.",
        ]),
        ("STEP 1 — PRE-PAYROLL DATA COLLECTION", [
            ("•", "Collect approved hours from client (timesheet, software, or verbal confirmation) by [Day] before payroll date"),
            ("•", "Confirm any changes: new hires, terminations, raises, bonuses, deductions"),
            ("•", "Verify PTO/sick balances if applicable"),
            ("•", "Check for garnishments, child support orders, or other court-ordered deductions"),
        ]),
        ("STEP 2 — PAYROLL PROCESSING", [
            ("•", "Enter hours and compensation changes into payroll software"),
            ("•", "Run payroll preview — review gross pay, deductions, net pay for each employee"),
            ("•", "Send preview to client for approval (require written/email sign-off)"),
            ("•", "Submit payroll only after client approval received"),
        ]),
        ("STEP 3 — TAX FILING", [
            ("•", "Payroll software handles federal tax deposits (FICA, FUTA, income tax withholding) automatically — verify timing"),
            ("•", "State payroll tax: confirm deposit schedule and deadlines by state"),
            ("•", "Quarterly: File Form 941 (federal) by April 30, July 31, Oct 31, Jan 31"),
            ("•", "Annually: File Form 940 (FUTA), W-2s (Jan 31), and state annual reconciliation"),
        ]),
        ("STEP 4 — POST-PAYROLL BOOKKEEPING", [
            ("•", "Record payroll journal entry in accounting software"),
            ("•", "Verify payroll bank account has sufficient funds before ACH hits"),
            ("•", "Reconcile payroll account monthly"),
            ("•", "File payroll reports and confirmations in client folder"),
        ]),
        ("PAYROLL TOOLS REFERENCE", [
            ("•", "Gusto: Best for small businesses; handles filing automatically"),
            ("•", "ADP: Enterprise-grade; complex setups"),
            ("•", "QuickBooks Payroll: Integrated with QBO; good for existing QBO clients"),
            ("•", "Paychex: Mid-market; strong compliance support"),
            ("•", "Wave Payroll: Affordable; limited to specific states"),
        ]),
    ])

doc("04_PAYROLL_SERVICES/Payroll_Compliance_Checklist.docx",
    "Payroll Compliance Checklist",
    "Bookkeeping Business Operating System | Payroll Services",
    [
        ("FEDERAL COMPLIANCE", [
            ("•", "[ ] Form W-4 on file for every employee"),
            ("•", "[ ] EIN (Employer Identification Number) obtained"),
            ("•", "[ ] Federal income tax withheld and deposited per schedule"),
            ("•", "[ ] FICA (Social Security + Medicare) withheld and matched by employer"),
            ("•", "[ ] FUTA (Federal Unemployment) calculated and paid"),
            ("•", "[ ] Form 941 filed quarterly"),
            ("•", "[ ] Form 940 filed annually"),
            ("•", "[ ] W-2s issued to all employees by January 31"),
            ("•", "[ ] W-3 filed with Social Security Administration by January 31"),
        ]),
        ("STATE COMPLIANCE", [
            ("•", "[ ] State income tax withheld (if applicable in state)"),
            ("•", "[ ] State unemployment (SUTA) registered and paid"),
            ("•", "[ ] State withholding account registered"),
            ("•", "[ ] State W-2 equivalent filed (if required)"),
            ("•", "[ ] New hire reporting submitted (usually within 20 days of hire)"),
        ]),
        ("EMPLOYEE RECORDS", [
            ("•", "[ ] Form I-9 completed and stored for all employees"),
            ("•", "[ ] Social Security numbers verified"),
            ("•", "[ ] Payroll records retained for minimum 3 years (federal) / 4 years (IRS)"),
            ("•", "[ ] Pay stubs provided each pay period (required in most states)"),
        ]),
        ("INDEPENDENT CONTRACTORS", [
            ("•", "[ ] Form W-9 collected before first payment"),
            ("•", "[ ] Form 1099-NEC issued by January 31 for payments of $600+"),
            ("•", "[ ] No employment taxes withheld for true contractors"),
            ("•", "[ ] Worker classification reviewed (misclassification risk assessed)"),
        ]),
        ("ANNUAL PAYROLL CALENDAR REMINDERS", [
            ("•", "January 31: W-2s to employees, W-3 to SSA, 1099s to contractors"),
            ("•", "April 30: Q1 Form 941 due"),
            ("•", "July 31: Q2 Form 941 due"),
            ("•", "October 31: Q3 Form 941 due"),
            ("•", "January 31 (following year): Q4 Form 941 and Form 940 due"),
        ]),
    ])

doc("04_PAYROLL_SERVICES/New_Employee_Setup_Checklist.docx",
    "New Employee Setup Checklist",
    "Bookkeeping Business Operating System | Payroll Services",
    [
        ("DOCUMENTS TO COLLECT", [
            ("•", "[ ] Signed offer letter"),
            ("•", "[ ] Completed Form W-4 (federal withholding)"),
            ("•", "[ ] State withholding form (if applicable)"),
            ("•", "[ ] Completed Form I-9 (with identity documents verified)"),
            ("•", "[ ] Direct deposit authorization form"),
            ("•", "[ ] Signed employee handbook acknowledgment"),
            ("•", "[ ] Benefits enrollment forms (health, dental, vision, 401k if offered)"),
        ]),
        ("SYSTEM SETUP", [
            ("•", "[ ] Add employee to payroll software"),
            ("•", "[ ] Enter: name, address, SSN, start date, pay rate, pay schedule"),
            ("•", "[ ] Set up withholding per W-4"),
            ("•", "[ ] Set up direct deposit banking information"),
            ("•", "[ ] Add any voluntary deductions (health insurance, 401k)"),
            ("•", "[ ] Set up PTO accrual policy if applicable"),
        ]),
        ("REPORTING", [
            ("•", "[ ] Submit new hire report to state (typically within 20 days)"),
            ("•", "[ ] Notify workers' compensation carrier of new hire"),
            ("•", "[ ] Add employee to benefits carrier portal"),
        ]),
        ("BOOKKEEPING SETUP", [
            ("•", "[ ] Add employee to payroll expense accounts"),
            ("•", "[ ] Update payroll projections in budget if applicable"),
            ("•", "[ ] Note employee start date for benefits eligibility tracking"),
        ]),
    ])

doc("04_PAYROLL_SERVICES/Payroll_Troubleshooting_Guide.docx",
    "Payroll Troubleshooting Guide",
    "Bookkeeping Business Operating System | Payroll Services",
    [
        ("ISSUE: EMPLOYEE PAID WRONG AMOUNT", [
            "Step 1: Identify the error — underpayment, overpayment, or incorrect deduction",
            "Underpayment: Issue a supplemental payroll run immediately for the difference",
            "Overpayment: Notify employee in writing; recover through deduction from next paycheck (check state law for limits)",
            "Tax impact: Supplemental wages are usually subject to flat federal rate (22% as of current law)",
        ]),
        ("ISSUE: PAYROLL TAX DEPOSIT MISSED", [
            "Step 1: Deposit immediately to minimize penalties",
            "Federal late deposit penalties: 2% (1–5 days late), 5% (6–15 days), 10% (16+ days), 15% (10+ days after first notice)",
            "Step 2: Determine if penalty abatement is possible (first-time penalty abatement from IRS)",
            "Step 3: Document the cause and implement process to prevent recurrence",
        ]),
        ("ISSUE: PAYROLL SOFTWARE CONNECTION FAILURE", [
            "Step 1: Check if bank account is connected and funded",
            "Step 2: Contact payroll software support immediately if ACH did not transmit",
            "Step 3: If payroll will be late, notify employees and client immediately",
            "Step 4: Process manual checks if necessary to meet state pay timing requirements",
        ]),
        ("ISSUE: W-2 DISCREPANCY AFTER YEAR-END", [
            "Step 1: Compare W-2 amounts to payroll register for the year",
            "Step 2: Identify timing differences, benefit amounts, or employer contributions",
            "Step 3: Issue corrected W-2c if error found; file with SSA and provide to employee",
            "Note: W-2c can be filed at any time but should be done as quickly as possible",
        ]),
        ("ISSUE: WORKER CLASSIFICATION QUESTION", [
            "The IRS uses a multi-factor test to distinguish employees from independent contractors.",
            ("•", "Behavioral control: Does the company control how, when, and where work is done? (Employee indicator)"),
            ("•", "Financial control: Does the company control business aspects like payment method? (Employee indicator)"),
            ("•", "Type of relationship: Written contract, benefits, permanent relationship?"),
            "If classification is uncertain: File Form SS-8 with IRS for determination. Misclassification can result in back taxes, penalties, and interest.",
        ]),
    ])

print("✓ 04_PAYROLL_SERVICES complete")

# ── 05_TAX_PREPARATION ───────────────────────────────────────────────────────
doc("05_TAX_PREPARATION/Tax_Season_SOP.docx",
    "Tax Season SOP",
    "Bookkeeping Business Operating System | Tax Preparation",
    [
        ("TAX SEASON TIMELINE (January–April)", [
            "January 1–15: Send document request checklist to all clients",
            "January 31: W-2s and 1099s available — verify clients received them",
            "February 1–28: Collect all client tax documents; complete year-end close for all clients",
            "March 1–31: Primary tax filing season — coordinate with CPAs; file simpler returns",
            "April 1–15: Final push before April 15 deadline; file extensions for complex returns",
        ]),
        ("YEAR-END CLOSE CHECKLIST", [
            ("•", "[ ] All bank accounts reconciled through December 31"),
            ("•", "[ ] All credit card accounts reconciled through December 31"),
            ("•", "[ ] Depreciation entries recorded (coordinate with CPA)"),
            ("•", "[ ] Prepaid expenses and accruals recorded"),
            ("•", "[ ] Owner's equity / draws reviewed and properly classified"),
            ("•", "[ ] Inventory count adjusted if applicable"),
            ("•", "[ ] 1099 vendor list confirmed and 1099s prepared"),
            ("•", "[ ] Final P&L and Balance Sheet run and reviewed"),
        ]),
        ("WORKING WITH THE CLIENT'S CPA", [
            ("•", "Send year-end financials to CPA by [date] in their preferred format (PDF + accountant access)"),
            ("•", "Provide supporting schedules: fixed asset list, loan amortization, officer compensation"),
            ("•", "Be available for CPA questions — respond within 24 hours during tax season"),
            ("•", "Review tax return once prepared — flag any discrepancies with your books"),
        ]),
        ("COMMON TAX SEASON ISSUES", [
            ("•", "Missing receipts for large expenses — request from client immediately"),
            ("•", "Undeposited sales (cash businesses) — reconcile POS reports to bank deposits"),
            ("•", "Shareholder/member loans: document properly to avoid reclassification as income"),
            ("•", "Vehicle use: ensure mileage log or actual expense records available"),
            ("•", "Home office: calculate square footage percentage if applicable"),
        ]),
    ])

doc("05_TAX_PREPARATION/Client_Tax_Document_Checklist.docx",
    "Client Tax Document Checklist",
    "Bookkeeping Business Operating System | Tax Preparation",
    [
        ("INCOME DOCUMENTS", [
            ("•", "[ ] Form W-2 from all employers (if owner also has employment income)"),
            ("•", "[ ] Form 1099-NEC (nonemployee compensation received)"),
            ("•", "[ ] Form 1099-MISC (rent, prizes, other income)"),
            ("•", "[ ] Form 1099-K (payment card and third-party network income — PayPal, Stripe, Square)"),
            ("•", "[ ] Form 1099-INT (bank interest income)"),
            ("•", "[ ] Form 1099-DIV (dividend income)"),
            ("•", "[ ] Any other income records not in accounting software"),
        ]),
        ("BUSINESS EXPENSE RECORDS", [
            ("•", "[ ] Receipts for major purchases (assets, equipment over $[X])"),
            ("•", "[ ] Vehicle mileage log or actual expense records"),
            ("•", "[ ] Home office measurements (if claiming deduction)"),
            ("•", "[ ] Travel expense records and business purpose documentation"),
            ("•", "[ ] Meal expense receipts with business purpose noted"),
        ]),
        ("LOAN AND ASSET DOCUMENTS", [
            ("•", "[ ] Year-end loan statements (balances as of December 31)"),
            ("•", "[ ] New equipment purchased (invoice + date placed in service)"),
            ("•", "[ ] Assets disposed of or sold during the year"),
            ("•", "[ ] Any real estate purchased or sold"),
        ]),
        ("PAYROLL DOCUMENTS", [
            ("•", "[ ] Year-end payroll summary (from payroll provider)"),
            ("•", "[ ] W-2s and W-3 filed with SSA"),
            ("•", "[ ] 1099s issued to contractors"),
            ("•", "[ ] Owner's health insurance premiums (if S-Corp)"),
            ("•", "[ ] Retirement contributions made (SEP-IRA, SIMPLE IRA, 401k)"),
        ]),
        ("PRIOR YEAR DOCUMENTS", [
            ("•", "[ ] Prior year tax return (all pages)"),
            ("•", "[ ] Prior year depreciation schedule (Form 4562)"),
            ("•", "[ ] Any IRS notices received during the year"),
        ]),
    ])

doc("05_TAX_PREPARATION/Quarterly_Tax_Estimates_Guide.docx",
    "Quarterly Tax Estimates Guide",
    "Bookkeeping Business Operating System | Tax Preparation",
    [
        ("WHO NEEDS TO PAY QUARTERLY ESTIMATES", [
            "Self-employed individuals and business owners who expect to owe $1,000 or more in federal taxes for the year typically need to pay quarterly estimated taxes. This includes sole proprietors, partners, S-corp shareholders, and single-member LLC owners.",
        ]),
        ("QUARTERLY DEADLINES", [
            ("•", "Q1 (January 1 – March 31): Due April 15"),
            ("•", "Q2 (April 1 – May 31): Due June 15"),
            ("•", "Q3 (June 1 – August 31): Due September 15"),
            ("•", "Q4 (September 1 – December 31): Due January 15 (following year)"),
        ]),
        ("CALCULATION METHOD 1 — SAFE HARBOR", [
            "Pay at least 100% of prior year tax liability in equal quarterly installments (110% if prior year AGI exceeded $150,000). This eliminates underpayment penalty even if actual tax is higher.",
            "Formula: Prior year tax / 4 = each quarterly payment",
        ]),
        ("CALCULATION METHOD 2 — ACTUAL ESTIMATE", [
            "Estimate current year income and expenses, calculate projected tax, subtract withholding:",
            "1. Run YTD P&L from accounting software",
            "2. Annualize the income (YTD / months elapsed × 12)",
            "3. Subtract estimated deductions",
            "4. Apply tax rates (federal + state + self-employment tax)",
            "5. Divide by 4 (or remaining quarters)",
        ]),
        ("HOW TO MAKE PAYMENTS", [
            ("•", "Federal: IRS Direct Pay (irs.gov/payments) — free; or IRS EFTPS"),
            ("•", "State: Each state has its own payment portal"),
            ("•", "Include: taxpayer name, SSN/EIN, tax year, payment type (1040-ES)"),
            ("•", "Keep confirmation numbers for every payment"),
        ]),
        ("BOOKKEEPING FOR TAX PAYMENTS", [
            ("•", "Record each payment as: Debit 'Owner's Draw' or 'Tax Payments'; Credit 'Bank'"),
            ("•", "Do NOT record as a business expense (personal income taxes are not deductible)"),
            ("•", "Track total payments per quarter in your client notes"),
        ]),
    ])

doc("05_TAX_PREPARATION/Year_End_Close_Checklist.docx",
    "Year-End Close Checklist",
    "Bookkeeping Business Operating System | Tax Preparation",
    [
        ("DECEMBER ACTIONS (Before Year-End)", [
            ("•", "[ ] Review fixed asset purchases for potential Section 179 deduction"),
            ("•", "[ ] Confirm retirement contributions to be made before December 31"),
            ("•", "[ ] Prepay deductible expenses if advantageous (rent, subscriptions)"),
            ("•", "[ ] Write off uncollectible receivables"),
            ("•", "[ ] Document charitable donations made"),
        ]),
        ("JANUARY ACTIONS (First 2 Weeks)", [
            ("•", "[ ] Send 1099-NEC to all qualifying contractors by January 31"),
            ("•", "[ ] File 1096 summary with IRS by January 31"),
            ("•", "[ ] Issue W-2s to all employees by January 31"),
            ("•", "[ ] File W-3 with Social Security Administration by January 31"),
            ("•", "[ ] Send tax document checklist to client"),
        ]),
        ("YEAR-END RECONCILIATION CHECKLIST", [
            ("•", "[ ] All bank accounts reconciled through December 31"),
            ("•", "[ ] All credit cards reconciled through December 31"),
            ("•", "[ ] Petty cash reconciled"),
            ("•", "[ ] Payroll liabilities cleared (no balance owed at year-end)"),
            ("•", "[ ] Sales tax liabilities cleared or properly accrued"),
            ("•", "[ ] Opening balance equity resolved to zero"),
            ("•", "[ ] Fixed assets list updated with additions and disposals"),
            ("•", "[ ] Depreciation recorded (coordinate with CPA for method)"),
        ]),
        ("FINAL REPORTS TO DELIVER TO CPA", [
            ("•", "[ ] Profit & Loss (full year, by month)"),
            ("•", "[ ] Balance Sheet (as of December 31)"),
            ("•", "[ ] Fixed asset schedule"),
            ("•", "[ ] Loan schedule"),
            ("•", "[ ] Payroll summary"),
            ("•", "[ ] 1099 vendor list"),
            ("•", "[ ] Any unusual items explained in writing"),
        ]),
    ])

print("✓ 05_TAX_PREPARATION complete")

# ── 06_CLIENT_COMMUNICATION ───────────────────────────────────────────────────
doc("06_CLIENT_COMMUNICATION/Monthly_Report_Email_Templates.docx",
    "Monthly Report Email Templates",
    "Bookkeeping Business Operating System | Client Communication",
    [
        ("STANDARD MONTHLY DELIVERY", [
            "Subject: Your [Month] Financial Reports — [Client Business Name]",
            "Hi [Name], your [Month] reports are attached. Here are the highlights:",
            ("•", "Revenue: $[X] ([+/-X%] vs. last month)"),
            ("•", "Expenses: $[X]"),
            ("•", "Net Profit: $[X] ([X]% margin)"),
            ("•", "Cash Balance: $[X] as of [date]"),
            "Notable this month: [1–2 sentence insight — e.g., 'Software subscriptions increased 18% — worth reviewing.']",
            "Action needed from you: [if any, or 'None this month — you're all set!']",
            "Next delivery: [date]. Questions? Reply anytime.\n\n[Your Name]",
        ]),
        ("STRONG MONTH", [
            "Subject: Great Month! Your [Month] Reports Are Ready",
            "Hi [Name], happy to share your [Month] reports — this was a strong one.",
            ("•", "Revenue: $[X] — your best month of the year so far!"),
            ("•", "Net Profit: $[X] ([X]% margin) — up from $[X] last month"),
            "Keep it going. Reports attached. Let me know if you'd like to dig into what drove the growth.",
        ]),
        ("SLOW MONTH WITH CONTEXT", [
            "Subject: [Month] Reports — A Softer Month + What to Watch",
            "Hi [Name], your [Month] reports are attached. A few notes:",
            ("•", "Revenue was $[X] — down [X]% from [Month], which is typical for this time of year."),
            ("•", "Net Profit: $[X]. Margins held steady at [X]%."),
            ("•", "Cash remains healthy at $[X] — no concerns there."),
            "No action needed — just keeping you informed. Always here if you have questions.",
        ]),
        ("DOCUMENT REQUEST", [
            "Subject: Quick Question About a [Month] Transaction",
            "Hi [Name], I have one quick question while closing your [Month] books:",
            "Transaction on [date]: $[X] from [vendor/description]. Can you clarify: [ ] Business expense — category: ___ [ ] Personal — I'll reclassify.",
            "No rush — just reply when you get a chance. I'll hold this item until I hear from you. Thanks!",
        ]),
    ])

doc("06_CLIENT_COMMUNICATION/Difficult_Conversation_Scripts.docx",
    "Difficult Conversation Scripts",
    "Bookkeeping Business Operating System | Client Communication",
    [
        ("SCENARIO: CLIENT IS SIGNIFICANTLY BEHIND ON PAYMENT", [
            "'[Name], I want to address something before we go further. Your account has an outstanding balance of $[X] covering [months]. I need to receive payment or a payment arrangement before continuing work. I value our relationship and want to keep your books current — what can we work out?'",
        ]),
        ("SCENARIO: BOOKS ARE IN WORSE SHAPE THAN EXPECTED", [
            "'[Name], now that I've had time to dig into your books, I want to be transparent. The catch-up work is more extensive than initially quoted — I'm seeing [X months] of unreconciled transactions and [specific issue]. To properly clean this up, I'll need to revise the catch-up fee to $[X]. I wanted to discuss this before proceeding. Does that work for you?'",
        ]),
        ("SCENARIO: CLIENT KEEPS SENDING INCOMPLETE INFORMATION", [
            "'[Name], I want to flag something that's been delaying your monthly close. I'm consistently missing [receipts / bank statements / expense reports] from you by my deadline. Without these, I can't produce accurate reports. Could we set up a system — [a shared folder / recurring reminder / specific day] — to make this easier for you?'",
        ]),
        ("SCENARIO: CLIENT QUESTIONS YOUR WORK", [
            "'[Name], I take your concern seriously. Let me review the [specific transaction / report / amount] and get back to you with an explanation. If there's an error on my end, I'll correct it immediately. Can you point me to exactly what you're looking at?'",
        ]),
        ("SCENARIO: RAISING FEES", [
            "'[Name], I'm sending a heads-up that my rates will increase starting [date]. Your new monthly rate will be $[X], up from $[X]. This reflects [scope expansion / cost increases / time invested]. I genuinely enjoy working with [Business] and want to continue our partnership. If you have questions, I'm happy to discuss.'",
        ]),
        ("SCENARIO: ENDING THE RELATIONSHIP", [
            "'[Name], after careful consideration, I've decided to discontinue our engagement effective [date]. This is not a reflection of your business — rather, [honest reason]. I'll ensure a smooth transition: your files, reports, and login credentials will be delivered within [10] business days. I wish you and [Business] continued success.'",
        ]),
    ])

doc("06_CLIENT_COMMUNICATION/Client_Questions_Scripts.docx",
    "Client Questions Scripts",
    "Bookkeeping Business Operating System | Client Communication",
    [
        ("QUESTION: WHAT'S MY PROFIT THIS MONTH?", [
            "'Your net profit for [month] was $[X]. That's your revenue of $[X] minus total expenses of $[X]. Your margin was [X]%, which [is above/below] your average for the year of [X]%.'",
        ]),
        ("QUESTION: CAN I AFFORD TO HIRE SOMEONE?", [
            "'Great question — let me pull your numbers. Your current monthly payroll is $[X] and your profit before payroll is averaging $[X]/month. Adding a [role] at $[X]/month would bring your payroll to $[X]. Based on current revenue, that's [sustainable / tight / I'd recommend reviewing in 2–3 months]. Want me to model out a few scenarios?'",
        ]),
        ("QUESTION: DO I OWE TAXES THIS QUARTER?", [
            "'Let me check your YTD numbers. Your net profit through [date] is $[X]. Based on your estimated tax rate of [X]%, you should have paid approximately $[X] in estimates this year. You've paid $[X], which means [you're on track / you may be underpaid by approximately $[X]]. I'd recommend running this by your CPA for a precise number.'",
        ]),
        ("QUESTION: WHY ARE MY EXPENSES SO HIGH THIS MONTH?", [
            "'Let me break it down. Your expenses for [month] were $[X], which is [X]% higher than last month. The biggest increases were: [Category 1] up $[X], [Category 2] up $[X]. The [Category 1] increase was due to [specific transaction]. Would you like to dig into any of these?'",
        ]),
        ("QUESTION: SHOULD I FORM AN S-CORP?", [
            "'That's a tax strategy question best answered by your CPA, who knows your full tax picture. What I can tell you is that the bookkeeping for an S-Corp is more complex — officer salary, distributions, and payroll are all separate. I'd be happy to share your current profit figures with your CPA so they can model the potential savings for you.'",
        ]),
    ])

doc("06_CLIENT_COMMUNICATION/Newsletter_Templates.docx",
    "Newsletter Templates",
    "Bookkeeping Business Operating System | Client Communication",
    [
        ("QUARTERLY NEWSLETTER TEMPLATE", [
            "Subject: [Bookkeeper Name] Quarterly Update — [Season] [Year]",
            "Hi [Name],",
            "Here's a quick roundup of what's new and what to watch as a small business owner this quarter.",
            "",
            "TAX REMINDER: [Specific quarterly deadline] is coming up on [date]. Here's what you need to know: [2–3 sentence relevant guidance].",
            "",
            "BOOKKEEPING TIP OF THE QUARTER: [One practical, actionable tip — e.g., 'Reconcile your books at least monthly — catching errors early is 10x faster than finding them at year-end.']",
            "",
            "WHAT'S NEW IN [ACCOUNTING SOFTWARE]: [One useful new feature or update relevant to clients]",
            "",
            "As always, reply to this email with any questions. Happy to help.",
            "[Your Name] | [Business Name] | [Phone]",
        ]),
        ("TAX SEASON NEWSLETTER", [
            "Subject: Tax Season Is Here — What You Need to Prepare",
            "Hi [Name], it's that time of year. Here's your quick tax season checklist:",
            ("•", "Documents you'll need: [W-2, 1099s, bank statements, receipts for large purchases]"),
            ("•", "What I'll handle: Year-end close, reports for your CPA, 1099s if applicable"),
            ("•", "Your action: Send me any documents not already in your accounting software"),
            ("•", "Deadline: I'll need everything by [date] to meet your CPA's deadline of [date]"),
            "Questions? Reply here or call [phone].",
        ]),
        ("NEW SERVICE ANNOUNCEMENT", [
            "Subject: Now Offering [New Service] — [Benefit Statement]",
            "Hi [Name], excited to share that I've expanded my services to include [new service].",
            "[2–3 sentences on what the service is and why it matters for clients like them.]",
            "If this sounds useful, reply to this email and I'll send you details. As a current client, you'll receive [priority access / discounted rate / first month free].",
        ]),
    ])

print("✓ 06_CLIENT_COMMUNICATION complete")

# ── 07_COMPLIANCE_LEGAL ───────────────────────────────────────────────────────
doc("07_COMPLIANCE_LEGAL/Confidentiality_Policy.docx",
    "Confidentiality Policy",
    "Bookkeeping Business Operating System | Compliance & Legal",
    [
        ("PURPOSE", [
            "[Business Name] handles sensitive financial information for multiple clients. This policy establishes our commitment to protecting client confidentiality and defines how financial data is handled.",
        ]),
        ("SCOPE", [
            "This policy applies to all employees, contractors, and partners of [Business Name] who have access to client financial information.",
        ]),
        ("DATA WE HANDLE", [
            ("•", "Business financial records (income, expenses, assets, liabilities)"),
            ("•", "Bank account numbers and routing numbers"),
            ("•", "Payroll data including employee names, compensation, and SSNs"),
            ("•", "Tax ID numbers (EIN, SSN)"),
            ("•", "Vendor and customer payment information"),
        ]),
        ("CONFIDENTIALITY OBLIGATIONS", [
            ("•", "All client financial data is strictly confidential"),
            ("•", "We never share, sell, or disclose client information to third parties without written consent"),
            ("•", "Exception: Disclosure required by law (IRS subpoena, court order)"),
            ("•", "Exception: Communication with client's designated CPA or tax professional (with client authorization)"),
        ]),
        ("DATA SECURITY PRACTICES", [
            ("•", "Client files stored in encrypted cloud storage with access controls"),
            ("•", "Accounting software access is password-protected with multi-factor authentication"),
            ("•", "No client data stored on personal devices without encryption"),
            ("•", "Client data is retained for [7] years after engagement ends, then securely deleted"),
        ]),
        ("BREACH NOTIFICATION", [
            "In the event of a data security breach affecting client information, [Business Name] will notify affected clients within [72] hours of discovery and describe the nature of the breach and steps taken to address it.",
        ]),
    ])

doc("07_COMPLIANCE_LEGAL/Data_Security_Policy.docx",
    "Data Security Policy",
    "Bookkeeping Business Operating System | Compliance & Legal",
    [
        ("CLOUD STORAGE AND FILE MANAGEMENT", [
            ("•", "All client files stored in [Google Drive / Dropbox / OneDrive] with 2FA enabled"),
            ("•", "Folder access limited to staff who work on that client"),
            ("•", "Client-specific folders shared only with client and their CPA (as authorized)"),
            ("•", "No files stored locally on personal devices"),
        ]),
        ("ACCOUNTING SOFTWARE ACCESS", [
            ("•", "Separate login credentials for each client's accounting software"),
            ("•", "Password manager required (LastPass, 1Password, or equivalent)"),
            ("•", "Multi-factor authentication enabled on all accounting platforms"),
            ("•", "Staff permissions set to 'accountant access' — never full administrator"),
        ]),
        ("EMAIL SECURITY", [
            ("•", "Client financial data transmitted only through encrypted email or secure portal"),
            ("•", "Never send SSNs, EINs, or account numbers via standard email — use secure file sharing"),
            ("•", "Password-protect PDFs containing sensitive data when emailed"),
        ]),
        ("DEVICE SECURITY", [
            ("•", "All devices used for client work must have screen lock and full disk encryption"),
            ("•", "Automatic screen lock after [5] minutes of inactivity"),
            ("•", "Software updates applied within [7] days of release"),
            ("•", "Antivirus/anti-malware on all devices"),
        ]),
        ("STAFF OFFBOARDING", [
            ("•", "Remove access to all client accounts within 24 hours of staff departure"),
            ("•", "Revoke email access, cloud storage, and accounting software access immediately"),
            ("•", "Change shared passwords that departing staff had access to"),
        ]),
    ])

doc("07_COMPLIANCE_LEGAL/Engagement_Letter_Template.docx",
    "Engagement Letter Template",
    "Bookkeeping Business Operating System | Compliance & Legal",
    [
        ("PURPOSE", [
            "An engagement letter is a shorter, more informal alternative to a full service agreement. It confirms the scope of work and terms. Use for smaller clients or simpler engagements.",
        ]),
        ("ENGAGEMENT LETTER", [
            "Date: [Date]",
            "To: [Client Name / Business Name]",
            "From: [Your Name / Business Name]",
            "",
            "Dear [Name],",
            "Thank you for choosing [Business Name] for your bookkeeping needs. This letter confirms the terms of our engagement.",
        ]),
        ("SCOPE OF SERVICES", [
            ("•", "Monthly bookkeeping for [Client Business Name]"),
            ("•", "Bank reconciliation for [X] accounts"),
            ("•", "Monthly Profit & Loss and Balance Sheet"),
            ("•", "Additional services: [list any add-ons]"),
            ("•", "Reports delivered by the [X]th of each month"),
        ]),
        ("FEES", [
            ("•", "Monthly fee: $[X], invoiced on the [1st] of each month"),
            ("•", "Payment due within [15] days"),
            ("•", "Late payment interest: [1.5]%/month after [30] days"),
        ]),
        ("TERMS", [
            ("•", "This engagement is month-to-month / [12-month] commitment"),
            ("•", "Either party may cancel with [30] days' written notice"),
            ("•", "This letter does not constitute tax or legal advice"),
        ]),
        ("AGREEMENT", [
            "Please sign below to confirm your acceptance of these terms.",
            "Client Signature: _________________________ Date: _________",
            "Your Signature: _________________________ Date: _________",
        ]),
    ])

print("✓ 07_COMPLIANCE_LEGAL complete")

# ── XLSX FILES ────────────────────────────────────────────────────────────────

# 1. Client Acquisition CRM
wb = Workbook()
ws = wb.active; ws.title = "Prospect Pipeline"
for col, w in zip("ABCDEFGH",[25,20,20,15,20,15,25,15]):
    ws.column_dimensions[col].width = w
hrow(ws,1,list(range(1,9)),["Business Name","Owner","Email","Phone","Stage","Est Monthly Fee","Last Touch","Next Action"])
rows=[
    ("Joe's Plumbing","Joe Morales","joe@joes.com","555-0101","Discovery Call","$400","Jan 10","Send proposal"),
    ("Riverdale Cafe","Maria Chen","maria@rcafe.com","555-0102","Proposal Sent","$600","Jan 8","Follow up"),
    ("Green Lawn LLC","Brian Park","brian@green.com","555-0103","Agreement","$350","Jan 5","Send agreement"),
    ("StyleHouse Salon","Amy Wu","amy@style.com","555-0104","Cold Outreach","$300","Jan 12","Day 3 follow-up"),
    ("BuildRight Co","Tom Davis","tom@build.com","555-0105","Negotiation","$800","Jan 3","Confirm scope"),
]
for i,r in enumerate(rows,2):
    drow(ws,i,list(range(1,9)),list(r),LGR if i%2==0 else WHT)
ws2=wb.create_sheet("Stage Legend")
hrow(ws2,1,[1,2],["Stage","Description"],ACC,WHT)
for i,r in enumerate([("Cold Outreach","Initial contact not yet made"),("Contacted","Message sent — awaiting response"),
    ("Discovery Call","Call scheduled/completed"),("Proposal Sent","Written proposal delivered"),
    ("Negotiation","Terms/price being discussed"),("Agreement","Engagement letter sent"),
    ("Active Client","Onboarded and providing services"),("Churned","Client ended relationship")],2):
    drow(ws2,i,[1,2],list(r),LGR if i%2==0 else WHT)
wb.save(BASE+"01_CLIENT_ACQUISITION/Client_Acquisition_CRM.xlsx")
print("  ✓ 01_CLIENT_ACQUISITION/Client_Acquisition_CRM.xlsx")

# 2. Client Onboarding Tracker
wb2=Workbook()
ws=wb2.active; ws.title="Onboarding Tracker"
for col,w in zip("ABCDEFG",[25,15,15,20,20,15,20]):
    ws.column_dimensions[col].width=w
hrow(ws,1,list(range(1,8)),["Client","Agreement Date","Kickoff Date","Software Access","First Report Date","Status","Monthly Fee"])
for i,r in enumerate([
    ("Joe's Plumbing","Jan 5","Jan 7","✓ Connected","Jan 20","In Progress","$400"),
    ("Riverdale Cafe","Jan 3","Jan 4","✓ Connected","Jan 15","First Report Sent","$600"),
    ("Green Lawn LLC","Dec 28","Dec 30","✓ Connected","Jan 10","Active","$350"),
    ("StyleHouse Salon","Jan 10","Jan 12","Awaiting Access","Jan 25","Awaiting Access","$300"),
],2):
    drow(ws,i,list(range(1,8)),list(r),LGR if i%2==0 else WHT)
ws2=wb2.create_sheet("Checklist")
hrow(ws2,1,[1,2,3],["Step","Action","Status"],ACC,WHT)
for i,r in enumerate([
    ("1","Agreement signed","✓"),("2","Welcome email sent","✓"),("3","Intake form completed","✓"),
    ("4","Software access granted","✓"),("5","Kickoff call completed","✓"),
    ("6","Chart of accounts reviewed","In Progress"),("7","Bank feeds connected","In Progress"),
    ("8","First reconciliation complete","Pending"),("9","First report delivered","Pending"),
    ("10","30-day check-in done","Pending")],2):
    drow(ws2,i,[1,2,3],list(r),LGR if i%2==0 else WHT)
wb2.save(BASE+"02_CLIENT_ONBOARDING/Client_Onboarding_Tracker.xlsx")
print("  ✓ 02_CLIENT_ONBOARDING/Client_Onboarding_Tracker.xlsx")

print("PART 2 DONE")
