"""Bookkeeping Business OS — Part 1: folders 00–03"""
import os, csv
from docx import Document
from docx.shared import Pt, RGBColor

BASE = "/home/user/oqul-phase55-production/bookkeeping-business-os/Ultimate_Bookkeeping_Business_Operating_System/"

FOLDERS = [
    "00_START_HERE","01_CLIENT_ACQUISITION","02_CLIENT_ONBOARDING",
    "03_BOOKKEEPING_SERVICES","04_PAYROLL_SERVICES","05_TAX_PREPARATION",
    "06_CLIENT_COMMUNICATION","07_COMPLIANCE_LEGAL","08_FINANCIAL_REPORTING",
    "09_BUSINESS_OPERATIONS","10_NOTION_WORKSPACE","11_BONUSES",
    "12_CANVA_IMPORTABLE_TEMPLATES"
]
for f in FOLDERS:
    os.makedirs(BASE + f, exist_ok=True)

def doc(filename, title, subtitle, sections):
    d = Document()
    t = d.add_paragraph(title); t.style = d.styles['Normal']
    t.runs[0].bold = True; t.runs[0].font.size = Pt(14)
    t.runs[0].font.color.rgb = RGBColor(0x0F,0x34,0x60)
    if subtitle:
        s = d.add_paragraph(subtitle); s.style = d.styles['Normal']
        s.runs[0].font.size = Pt(10)
        s.runs[0].font.color.rgb = RGBColor(0x7F,0x8C,0x8D)
    for section in sections:
        if isinstance(section, str):
            d.add_paragraph(section); continue
        heading, items = section
        h = d.add_heading(heading, level=1)
        h.runs[0].font.color.rgb = RGBColor(0x0F,0x34,0x60)
        for item in items:
            if isinstance(item, tuple) and item[0] == '•':
                d.add_paragraph(item[1], style='List Bullet')
            else:
                d.add_paragraph(str(item))
    d.save(BASE + filename)
    print(f"  ✓ {filename}")

# ── 00_START_HERE ──────────────────────────────────────────────────────────────
with open(BASE + "00_START_HERE/README_FIRST.txt", "w") as f:
    f.write("""ULTIMATE BOOKKEEPING BUSINESS OPERATING SYSTEM

Start with START_HERE_Implementation_Guide.pdf.

EDITABLE FORMATS
- DOCX: edit in Microsoft Word or upload to Google Docs.
- XLSX/CSV: edit in Excel or upload to Google Sheets.
- PPTX: upload to Canva to convert into editable Canva designs.
- Notion assets: import CSV databases and copy the workspace structure.

IMPORTANT
Contracts, policies, and agreements are generic educational templates and are not legal advice.
Have a qualified professional review them before commercial use.
Tax and financial guidance in this system is for educational purposes only.
Consult a licensed CPA or tax professional for jurisdiction-specific advice.
""")
print("  ✓ 00_START_HERE/README_FIRST.txt")

# ── 01_CLIENT_ACQUISITION ─────────────────────────────────────────────────────
doc("01_CLIENT_ACQUISITION/Cold_Outreach_Templates.docx",
    "Cold Outreach Templates",
    "Bookkeeping Business Operating System | Client Acquisition",
    [
        ("LINKEDIN CONNECTION REQUEST TEMPLATES", [
            ("•", "Small Business Owner: 'Hi [Name], I help [industry] business owners keep their books clean and their taxes low. I'd love to connect — I share useful bookkeeping tips here regularly.'"),
            ("•", "Restaurant Owner: 'Hi [Name], I specialize in bookkeeping for restaurants and food businesses. Noticed you run [Restaurant] — I'd love to connect and share some cost-saving insights.'"),
            ("•", "E-commerce Seller: 'Hi [Name], I help e-commerce sellers on [Amazon/Shopify/Etsy] understand their true profit margins. Happy to connect!'"),
            ("•", "Construction Contractor: 'Hi [Name], job costing and contractor bookkeeping is a specialty of mine. Always happy to connect with local business owners.'"),
        ]),
        ("COLD EMAIL TEMPLATES", [
            ("•", "Subject: Are Your Books Ready for Tax Season?\nHi [Name], tax season is [X months] away. Many [industry] business owners I work with were surprised by their tax bill because their books weren't current. I specialize in monthly bookkeeping for [industry] businesses — keeping you accurate, organized, and tax-ready year-round. Worth a 15-minute call? — [Your Name]"),
            ("•", "Subject: How Much Did [Business Name] Profit Last Month?\nHi [Name], if you can't answer that question immediately, you're not alone — most small business owners can't. I provide clean monthly financials so you always know where you stand. Happy to share how. — [Your Name]"),
            ("•", "Subject: $[X] Saved for a [Industry] Client Last Year\nHi [Name], I caught $[X] in deductions for a [similar business] client last year that their previous bookkeeper missed. I specialize in [industry] bookkeeping. 15 minutes to see if I can help you too? — [Your Name]"),
        ]),
        ("PHONE COLD CALL SCRIPT", [
            "Opening: 'Hi [Name], this is [Your Name] from [Business Name]. I specialize in bookkeeping for [industry] businesses. Do you have 2 minutes?'",
            "Value statement: 'I help [industry] owners get clean monthly financials, stay tax-ready, and actually understand their numbers. Most of my clients save [X] hours a month and reduce their tax bill.'",
            "Question: 'Are you currently working with a bookkeeper, or handling your own books?'",
            "If DIY: 'How many hours a month do you spend on bookkeeping? I can probably take that off your plate entirely.'",
            "Close: 'Could we schedule a 15-minute call this week to see if there's a fit?'",
        ]),
        ("FOLLOW-UP SEQUENCE (After No Response)", [
            ("•", "Day 3: 'Hi [Name], following up on my message about bookkeeping support for [Business]. Happy to share a quick overview of how I work.'"),
            ("•", "Day 7: 'Hi [Name], I know tax season sneaks up fast. I have one open slot for a new monthly client this quarter. Want to grab it?'"),
            ("•", "Day 14: 'Hi [Name], last follow-up. If timing isn't right, I'd love to reconnect before tax season. Here's my calendar: [link]'"),
        ]),
        ("REFERRAL OUTREACH", [
            ("•", "To CPA/accountant: 'Hi [Name], I specialize in bookkeeping and often get clients who need a CPA for taxes. I'd love to establish a referral relationship — I send you tax clients; you send me bookkeeping clients. Open to a quick call?'"),
            ("•", "To existing client: '[Name], I really enjoy working with [Business]. If you know any business owners who are struggling with their books, I'd love an introduction. I'll make sure they're well taken care of.'"),
        ]),
    ])

doc("01_CLIENT_ACQUISITION/Discovery_Call_Script.docx",
    "Discovery Call Script",
    "Bookkeeping Business Operating System | Client Acquisition",
    [
        ("OPENING (0–2 min)", [
            "'[Name], thanks for connecting. My goal today is to understand your business and see if I can genuinely help. I'll ask a few questions and you can do the same. Sound good?'",
        ]),
        ("SITUATION QUESTIONS (2–10 min)", [
            ("•", "Tell me about your business — how long have you been operating, and how many employees?"),
            ("•", "How are you currently handling your bookkeeping?"),
            ("•", "How often are your books fully up to date?"),
            ("•", "What accounting software are you using, if any?"),
            ("•", "Do you have a CPA or tax preparer you work with?"),
        ]),
        ("PAIN IDENTIFICATION (10–18 min)", [
            ("•", "What's the most frustrating part of managing your finances right now?"),
            ("•", "Have you ever had a surprise tax bill — something higher than you expected?"),
            ("•", "Do you feel confident you're capturing all your deductions?"),
            ("•", "What happens to your business if you're not clear on your cash position each month?"),
            ("•", "How many hours a month do you spend on bookkeeping tasks?"),
        ]),
        ("GOALS & SCOPE (18–22 min)", [
            ("•", "What would 'perfect' look like for your books?"),
            ("•", "Do you need payroll processing as well?"),
            ("•", "Are you looking for help with sales tax / quarterly estimates?"),
            ("•", "What accounting software would you prefer (QuickBooks, Xero, Wave)?"),
        ]),
        ("PRESENTING YOUR SOLUTION (22–27 min)", [
            "Based on what you've shared, here's how I'd approach this: Monthly bookkeeping includes [scope], delivered by the [X]th of each month. You'd receive a Profit & Loss, Balance Sheet, and a plain-English summary of what the numbers mean. I also flag anything that looks unusual.",
            "My monthly fee for your scope is $[X]. That includes [services]. There's no contract — you can cancel with 30 days' notice.",
        ]),
        ("CLOSE (27–30 min)", [
            ("•", "Does that sound like what you're looking for?"),
            ("•", "If I can get your books current within [X] weeks and deliver your first report by [date], would you be ready to start?"),
            ("•", "The next step is signing my engagement letter and completing a short intake form. Can we do that this week?"),
        ]),
    ])

doc("01_CLIENT_ACQUISITION/Service_Proposal_Template.docx",
    "Service Proposal Template",
    "Bookkeeping Business Operating System | Client Acquisition",
    [
        ("EXECUTIVE SUMMARY", [
            "[Your Business Name] is pleased to present this proposal to [Client Business Name] for bookkeeping services. Based on our conversation on [Date], we understand your key needs are [brief summary].",
        ]),
        ("SERVICES INCLUDED", [
            ("•", "Monthly bank and credit card reconciliation (all accounts)"),
            ("•", "Categorization of all income and expenses"),
            ("•", "Monthly Profit & Loss Statement"),
            ("•", "Monthly Balance Sheet"),
            ("•", "Accounts payable tracking (if applicable)"),
            ("•", "Accounts receivable tracking (if applicable)"),
            ("•", "Monthly financial summary email with plain-English insights"),
            ("•", "Coordination with your CPA / tax preparer at year-end"),
        ]),
        ("ADD-ON SERVICES (Available Separately)", [
            ("•", "Payroll processing: $[X]/month"),
            ("•", "Sales tax filing: $[X]/filing"),
            ("•", "Quarterly estimated tax calculations: $[X]/quarter"),
            ("•", "Catch-up bookkeeping: $[X]/month (for months behind)"),
            ("•", "Financial analysis and forecasting: $[X]/month"),
        ]),
        ("INVESTMENT", [
            ("•", "Monthly bookkeeping: $[X]/month"),
            ("•", "One-time setup fee: $[X] (waived if annual commitment)"),
            ("•", "Catch-up bookkeeping (if applicable): $[X] — covers [date range]"),
            ("•", "Payment: Invoice sent on the [1st] of each month; due within [15] days"),
            ("•", "Contract: [Month-to-month / 12-month commitment] with [30]-day written cancellation notice"),
        ]),
        ("WHY [YOUR BUSINESS NAME]", [
            ("•", "Specialization in [industry/niche] businesses"),
            ("•", "Certified in [QuickBooks / Xero / Wave]"),
            ("•", "[X]+ clients served across [industry]"),
            ("•", "Clean books delivered by the [X]th of every month — guaranteed"),
            ("•", "Direct communication — no offshore team, no outsourcing"),
        ]),
        ("NEXT STEPS", [
            "1. Review and sign the Engagement Letter",
            "2. Complete the Client Intake Form (10 minutes)",
            "3. Grant access to your accounting software and bank feeds",
            "4. Kickoff call to confirm scope and first delivery date",
        ]),
    ])

doc("01_CLIENT_ACQUISITION/Follow_Up_Sequence.docx",
    "Follow-Up Sequence",
    "Bookkeeping Business Operating System | Client Acquisition",
    [
        ("7-TOUCH FOLLOW-UP FRAMEWORK", [
            "Most clients need 4–7 touches before deciding. Use this sequence consistently across cold outreach and warm leads.",
        ]),
        ("TOUCH 1 — Day 1 (Initial Outreach)", [
            ("•", "Channel: Email or LinkedIn"),
            ("•", "Goal: Start the conversation"),
            ("•", "Message: [See Cold Outreach Templates]"),
        ]),
        ("TOUCH 2 — Day 3 (Value Add)", [
            ("•", "Channel: Same as Touch 1"),
            ("•", "Message: 'Hi [Name], shared a tip I thought was relevant for [industry] owners — [1 bookkeeping insight, e.g., how to track mileage deductions]. No action needed — just thought it might save you money.'"),
        ]),
        ("TOUCH 3 — Day 7 (Social Proof)", [
            ("•", "Channel: Email"),
            ("•", "Message: 'Hi [Name], I recently helped a [similar business] owner who hadn't reconciled their accounts in 8 months. We caught [3] uncategorized expenses totaling $[X] in missed deductions. 15-minute call?'"),
        ]),
        ("TOUCH 4 — Day 14 (Direct Ask)", [
            ("•", "Channel: Phone call"),
            ("•", "Script: 'Hi [Name], I've reached out a couple of times about bookkeeping support for [Business]. I don't want to keep bothering you — is this something worth discussing, or should I circle back at a different time?'"),
        ]),
        ("TOUCH 5 — Day 30 (Seasonal Hook)", [
            ("•", "Message: 'Hi [Name], [quarter/tax season] is approaching. A lot of business owners are scrambling right now. I have one open slot for a new client this month — thought of you first.'"),
        ]),
        ("TOUCH 6 — Day 60 (New Value)", [
            ("•", "Message: 'Hi [Name], I published a quick guide on the [X] most commonly missed deductions for [industry] businesses. Thought you might find it useful: [link or attach]. No commitment — just sharing.'"),
        ]),
        ("TOUCH 7 — Day 90 (Final Check-In)", [
            ("•", "Message: 'Hi [Name], last follow-up from me. If you ever need clean books, reliable monthly reports, or just want someone to take bookkeeping off your plate — I'm here. Here's my calendar: [link]'"),
        ]),
    ])

doc("01_CLIENT_ACQUISITION/Objection_Handling_Scripts.docx",
    "Objection Handling Scripts",
    "Bookkeeping Business Operating System | Client Acquisition",
    [
        ("OBJECTION: 'I DO MY OWN BOOKS'", [
            "Acknowledge: 'That's great — a lot of business owners start that way.'",
            "Reframe: 'Can I ask how many hours a month it takes you? And how confident are you that you're capturing every deduction?'",
            "Opportunity: 'Most of my clients were doing their own books before coming to me. On average they reclaim 6–10 hours a month and reduce their tax bill because a professional eye catches things they didn't know to look for.'",
        ]),
        ("OBJECTION: 'MY ACCOUNTANT HANDLES IT'", [
            "Clarify: 'When you say accountant — do they do monthly bookkeeping, or primarily your annual tax return?'",
            "Educate: 'Most CPAs focus on tax preparation, not ongoing bookkeeping. They typically get your records once a year. Monthly bookkeeping is a different service — it's the ongoing record-keeping that makes tax season fast and inexpensive.'",
            "Position: 'Many of my clients work with a CPA for taxes and me for monthly books. I can coordinate with them directly.'",
        ]),
        ("OBJECTION: 'YOUR PRICE IS TOO HIGH'", [
            "Reframe: 'I understand. Let me ask — how much do you pay your CPA to untangle a year of messy books at tax time? Many clients pay $[X]–$[X] in extra CPA fees because their records aren't clean. My monthly fee often pays for itself.'",
            "Tier down: 'If the full scope is too much right now, we could start with [smaller service] at $[X]/month and add services as you grow.'",
        ]),
        ("OBJECTION: 'I'M NOT READY YET / MAYBE NEXT QUARTER'", [
            "Acknowledge: 'That's fair. Can I ask what would need to change for you to feel ready?'",
            "Plant urgency: 'One thing to consider — the longer books go unreconciled, the longer and more expensive it is to catch up. Starting now means your Q[X] numbers will be clean for tax filing.'",
            "Schedule: 'Could I follow up with you in [30] days? I'll put a reminder on my end.'",
        ]),
        ("OBJECTION: 'WE HAVE INTERNAL STAFF FOR THAT'", [
            "Clarify: 'Great — are they a part-time bookkeeper or an office manager who also does books?'",
            "Complement: 'I often work alongside internal staff — they handle day-to-day data entry and I handle the monthly close, reconciliations, and financial reporting. This model often reduces errors and keeps reports consistent.'",
        ]),
    ])

doc("01_CLIENT_ACQUISITION/Pricing_Guide.docx",
    "Pricing Guide",
    "Bookkeeping Business Operating System | Client Acquisition",
    [
        ("PRICING PHILOSOPHY", [
            "Price based on value, not hours. Your clients care about clean books, tax savings, and financial clarity — not how long it takes you. Pricing by scope (transactions + accounts + services) is more predictable and scalable than hourly billing.",
        ]),
        ("PRICING BY TRANSACTION VOLUME (Monthly Bookkeeping)", [
            ("•", "Starter (0–100 transactions/month): $[200–400]/month — 1–2 bank accounts, P&L + Balance Sheet"),
            ("•", "Growth (101–300 transactions/month): $[400–750]/month — 2–4 accounts, all standard reports"),
            ("•", "Established (301–600 transactions/month): $[750–1,200]/month — 4–6 accounts, payroll coordination"),
            ("•", "Enterprise (600+ transactions/month): $[1,200+]/month — custom scope, multiple entities"),
        ]),
        ("ADD-ON SERVICE PRICING", [
            ("•", "Payroll processing: $[50–150]/month for 1–5 employees; $[15–25] per additional employee"),
            ("•", "Sales tax filing: $[75–150]/filing per state"),
            ("•", "Quarterly estimated tax calculations: $[100–200]/quarter"),
            ("•", "1099 preparation: $[50]/form + $[100] base fee"),
            ("•", "Catch-up bookkeeping: [1.5–2x] your standard monthly rate per month of backlog"),
            ("•", "CFO / advisory services: $[500–2,000]/month"),
            ("•", "Year-end close and tax prep coordination: $[200–500]/year"),
        ]),
        ("PRICING FACTORS THAT JUSTIFY HIGHER RATES", [
            ("•", "Industry specialization (restaurants, contractors, e-commerce — premium 20–30%)"),
            ("•", "Multiple entities or locations"),
            ("•", "Inventory tracking required"),
            ("•", "Job costing required (construction, project-based businesses)"),
            ("•", "Catch-up needed before ongoing service begins"),
            ("•", "Urgent turnaround or special reporting requirements"),
        ]),
        ("HOW TO PRESENT PRICING", [
            ("•", "Always present 3 tiers: Good / Better / Best"),
            ("•", "Anchor high: present your highest package first"),
            ("•", "Emphasize ROI, not cost: 'For $[X]/month, you get [specific outcomes]'"),
            ("•", "Offer annual pricing with a discount: 'Pay 12 months upfront and save 10%'"),
        ]),
    ])

print("✓ 01_CLIENT_ACQUISITION complete")

# ── 02_CLIENT_ONBOARDING ──────────────────────────────────────────────────────
doc("02_CLIENT_ONBOARDING/Welcome_Email_Templates.docx",
    "Welcome Email Templates",
    "Bookkeeping Business Operating System | Client Onboarding",
    [
        ("EMAIL 1 — ENGAGEMENT LETTER SIGNED", [
            "Subject: Welcome to [Business Name] — Let's Get Your Books Set Up",
            "Hi [Name], welcome aboard! I'm excited to work with [Client Business]. Here's what happens next:",
            ("•", "Step 1: Complete the Client Intake Form — [link] (10 minutes)"),
            ("•", "Step 2: Grant access to your accounting software and bank feeds — instructions in the intake form"),
            ("•", "Step 3: We schedule a 20-minute kickoff call to confirm scope and first delivery date"),
            "Your first set of financials will be delivered by [date]. Questions? Reply here or call [phone]. Talk soon,\n[Your Name]",
        ]),
        ("EMAIL 2 — ACCESS GRANTED / SETUP CONFIRMED", [
            "Subject: Got Access — Starting Your Bookkeeping Now",
            "Hi [Name], I've connected to [QuickBooks/Xero/Wave] and your bank feeds. Everything looks good.",
            "Here's my plan for the first few weeks:",
            ("•", "Week 1: Review chart of accounts and clean up any miscategorizations"),
            ("•", "Week 2: Complete [Month/Year] reconciliations"),
            ("•", "Week 3: Deliver your first monthly reports"),
            "If anything looks unusual, I'll flag it with a quick email before making changes. Talk soon,\n[Your Name]",
        ]),
        ("EMAIL 3 — FIRST REPORT DELIVERED", [
            "Subject: Your [Month] Financial Reports Are Ready",
            "Hi [Name], your [Month] financial reports are attached:",
            ("•", "Profit & Loss Statement"),
            ("•", "Balance Sheet"),
            ("•", "Monthly Summary — plain-English notes on what the numbers mean"),
            "A few highlights from this month: [2–3 bullet insights]",
            "Questions? Reply to this email or let's hop on a quick call. I'll be back in touch by [date] with [Month+1] reports.",
        ]),
        ("EMAIL 4 — MONTHLY REPORT STANDARD DELIVERY", [
            "Subject: [Month] Financial Reports — [Client Business Name]",
            "Hi [Name], your [Month] reports are attached.",
            ("•", "Revenue: $[X] ([vs last month / vs budget])"),
            ("•", "Expenses: $[X]"),
            ("•", "Net Profit: $[X] ([margin]%)"),
            ("•", "Notable items: [brief note]"),
            ("•", "Action needed from you: [if any]",),
            "Let me know if you have any questions. Next delivery: [date].",
        ]),
    ])

doc("02_CLIENT_ONBOARDING/Client_Intake_Form.docx",
    "Client Intake Form",
    "Bookkeeping Business Operating System | Client Onboarding",
    [
        ("BUSINESS INFORMATION", [
            ("•", "Business Legal Name: _______________________________"),
            ("•", "DBA (if different): _______________________________"),
            ("•", "Business Entity Type: [ ] Sole Prop  [ ] LLC  [ ] S-Corp  [ ] C-Corp  [ ] Partnership"),
            ("•", "EIN / Tax ID: _______________________________"),
            ("•", "State of Formation: _______________________________"),
            ("•", "Business Start Date: _______________________________"),
            ("•", "Primary Industry: _______________________________"),
        ]),
        ("OWNER / CONTACT INFORMATION", [
            ("•", "Owner Name(s): _______________________________"),
            ("•", "Email: _______________________________"),
            ("•", "Phone: _______________________________"),
            ("•", "Mailing Address: _______________________________"),
        ]),
        ("ACCOUNTING & BANKING", [
            ("•", "Current accounting software: [ ] QuickBooks Online  [ ] Xero  [ ] Wave  [ ] None  [ ] Other: ___"),
            ("•", "Number of bank accounts: _______"),
            ("•", "Number of credit card accounts: _______"),
            ("•", "Do you accept payments via: [ ] PayPal  [ ] Stripe  [ ] Square  [ ] Venmo  [ ] Other: ___"),
            ("•", "Do you use a payroll service? [ ] Yes — Provider: ___  [ ] No"),
        ]),
        ("CURRENT BOOKS STATUS", [
            ("•", "Are your books current? [ ] Yes, within 30 days  [ ] 1–3 months behind  [ ] 3+ months behind"),
            ("•", "Last reconciliation date: _______________________________"),
            ("•", "Do you have a prior bookkeeper's files to transfer? [ ] Yes  [ ] No"),
            ("•", "Fiscal year end: [ ] December 31  [ ] Other: ___"),
        ]),
        ("YOUR CPA / TAX PREPARER", [
            ("•", "CPA / Tax Preparer Name: _______________________________"),
            ("•", "Firm: _______________________________"),
            ("•", "Email: _______________________________"),
            ("•", "Phone: _______________________________"),
            ("•", "May I contact them directly? [ ] Yes  [ ] No"),
        ]),
        ("SERVICES REQUESTED", [
            ("•", "[ ] Monthly bookkeeping"),
            ("•", "[ ] Payroll processing — Number of employees: ___"),
            ("•", "[ ] Sales tax filing — States: ___"),
            ("•", "[ ] Quarterly estimated tax calculations"),
            ("•", "[ ] Catch-up bookkeeping — Months behind: ___"),
            ("•", "[ ] 1099 preparation"),
            ("•", "[ ] Financial reporting / dashboards"),
        ]),
        ("ADDITIONAL NOTES", [
            ("•", "Anything else I should know about your business or finances? _______________________________"),
        ]),
    ])

doc("02_CLIENT_ONBOARDING/Kickoff_Call_Script.docx",
    "Kickoff Call Script",
    "Bookkeeping Business Operating System | Client Onboarding",
    [
        ("BEFORE THE CALL", [
            ("•", "Review the completed intake form"),
            ("•", "Log in to their accounting software — identify any obvious issues"),
            ("•", "Check bank connections are active"),
            ("•", "Note any discrepancies between what was discussed and intake form"),
        ]),
        ("OPENING (0–2 min)", [
            "'[Name], thanks for getting the intake form back. I've had a chance to review it and get access to [software]. The goal today is to align on scope, discuss first steps, and confirm your first delivery date. Should take about 20 minutes.'",
        ]),
        ("REVIEW SCOPE (2–8 min)", [
            ("•", "Confirm: accounts to reconcile, payroll, sales tax, catch-up needs"),
            ("•", "Ask: 'Are there any transactions or accounts I should know about before we start?'"),
            ("•", "Clarify: 'How do you handle [cash/tips/inventory/owner draws]?'"),
        ]),
        ("CHART OF ACCOUNTS ALIGNMENT (8–13 min)", [
            ("•", "'I'll review your chart of accounts and suggest any adjustments to match your industry. I'll email you before making any changes.'"),
            ("•", "'What categories are most important for you to track separately?'"),
            ("•", "'Do you need job costing by project/client?'"),
        ]),
        ("COMMUNICATION & DELIVERABLES (13–18 min)", [
            "'Here's how we'll work together: I deliver your monthly reports by the [X]th of each month. I'll email you if I have a question — I'll try to batch questions so I'm not in your inbox constantly. What's the best way to reach you for quick questions?'",
            "'If something looks unusual in your books, I'll flag it for you before categorizing. Does that work?'",
        ]),
        ("CLOSE (18–20 min)", [
            "'Any questions before I get started? Great. I'll have your first set of reports to you by [date]. I'll send a quick email if I hit anything unexpected.'",
        ]),
    ])

doc("02_CLIENT_ONBOARDING/Service_Agreement_Template.docx",
    "Service Agreement Template",
    "Bookkeeping Business Operating System | Client Onboarding",
    [
        ("PARTIES", [
            "This Service Agreement ('Agreement') is between [Your Business Name] ('Bookkeeper') and [Client Business Name] ('Client'), effective [Date].",
        ]),
        ("SERVICES", [
            "Bookkeeper agrees to provide the following services as described in the attached Scope of Services:",
            ("•", "Monthly bookkeeping and bank reconciliation"),
            ("•", "Monthly financial reports (P&L and Balance Sheet)"),
            ("•", "Additional services as specified in Exhibit A"),
        ]),
        ("FEES AND PAYMENT", [
            ("•", "Monthly fee: $[X], invoiced on the [1st] of each month"),
            ("•", "Payment due within [15] days of invoice"),
            ("•", "Late payment: [1.5]% per month on balances over [30] days"),
            ("•", "Setup / catch-up fees: as quoted separately"),
        ]),
        ("CLIENT RESPONSIBILITIES", [
            ("•", "Provide access to accounting software, bank feeds, and documents within [5] business days of request"),
            ("•", "Review and respond to bookkeeper's questions within [5] business days"),
            ("•", "Provide any additional information needed for accurate record-keeping"),
            ("•", "Notify bookkeeper of significant business changes (new accounts, major purchases, entity changes)"),
        ]),
        ("CONFIDENTIALITY", [
            "Both parties agree to maintain the confidentiality of all financial and business information shared. Bookkeeper will not disclose Client's financial information to any third party without prior written consent, except as required by law.",
        ]),
        ("LIMITATION OF LIABILITY", [
            "Bookkeeper's liability is limited to the fees paid in the [3] months preceding a claim. Bookkeeper is not liable for tax penalties, interest, or losses resulting from inaccurate information provided by Client. This agreement does not constitute tax, legal, or financial advice.",
        ]),
        ("TERMINATION", [
            "Either party may terminate with [30] days' written notice. Client remains responsible for all fees incurred through the termination date. Bookkeeper will provide all files and records to Client upon request within [10] business days of termination.",
        ]),
        ("GOVERNING LAW", [
            "This Agreement is governed by the laws of [State]. Disputes shall be resolved in [State] courts.",
        ]),
        ("SIGNATURES", [
            "Bookkeeper: _________________________ Date: _________",
            "Client: _________________________ Date: _________",
        ]),
    ])

doc("02_CLIENT_ONBOARDING/Onboarding_Checklist_SOP.docx",
    "Onboarding Checklist SOP",
    "Bookkeeping Business Operating System | Client Onboarding",
    [
        ("DAY 1 — AGREEMENT SIGNED", [
            ("•", "[ ] Send welcome email with intake form link"),
            ("•", "[ ] Create client record in practice management system"),
            ("•", "[ ] Set up client folder (cloud storage)"),
            ("•", "[ ] Schedule kickoff call within 3–5 days"),
        ]),
        ("DAYS 2–3 — INTAKE FORM RECEIVED", [
            ("•", "[ ] Review intake form for completeness"),
            ("•", "[ ] Request access to accounting software"),
            ("•", "[ ] Request access to bank/credit card feeds"),
            ("•", "[ ] Request prior bookkeeper files if applicable"),
            ("•", "[ ] Note catch-up scope and quote if needed"),
        ]),
        ("DAYS 3–5 — KICKOFF CALL", [
            ("•", "[ ] Log in to accounting software before call"),
            ("•", "[ ] Complete kickoff call using Kickoff Call Script"),
            ("•", "[ ] Confirm services, delivery date, and communication method"),
            ("•", "[ ] Send post-kickoff email summary"),
        ]),
        ("WEEK 1 — INITIAL SETUP", [
            ("•", "[ ] Review and clean up chart of accounts"),
            ("•", "[ ] Connect all bank and credit card feeds"),
            ("•", "[ ] Review prior period transactions for context"),
            ("•", "[ ] Document client-specific categorization notes"),
            ("•", "[ ] Begin current month reconciliation"),
        ]),
        ("FIRST REPORT DELIVERY", [
            ("•", "[ ] Complete all reconciliations"),
            ("•", "[ ] Run P&L and Balance Sheet"),
            ("•", "[ ] Write monthly summary notes"),
            ("•", "[ ] Send first report email using template"),
            ("•", "[ ] Schedule 15-minute review call if desired"),
        ]),
    ])

print("✓ 02_CLIENT_ONBOARDING complete")

# ── 03_BOOKKEEPING_SERVICES ───────────────────────────────────────────────────
doc("03_BOOKKEEPING_SERVICES/Monthly_Bookkeeping_SOP.docx",
    "Monthly Bookkeeping SOP",
    "Bookkeeping Business Operating System | Bookkeeping Services",
    [
        ("MONTHLY CLOSE TIMELINE", [
            ("•", "Day 1–3 of month: Download prior month bank and credit card statements"),
            ("•", "Day 1–5: Import transactions into accounting software"),
            ("•", "Day 3–8: Categorize and reconcile all accounts"),
            ("•", "Day 8–12: Review for unusual items; send questions to client"),
            ("•", "Day 12–15: Finalize reconciliations; run financial reports"),
            ("•", "By [X]th: Deliver reports to client with summary email"),
        ]),
        ("STEP 1 — TRANSACTION IMPORT AND CATEGORIZATION", [
            ("•", "Connect bank feeds or manually import CSV statements"),
            ("•", "Categorize all transactions using client's chart of accounts"),
            ("•", "Flag any transaction over $[X] for client confirmation"),
            ("•", "Identify and record recurring transactions"),
            ("•", "Separate personal vs. business transactions (document mixed items)"),
        ]),
        ("STEP 2 — BANK RECONCILIATION", [
            ("•", "Compare bank statement ending balance to accounting software"),
            ("•", "Identify any missing or duplicate transactions"),
            ("•", "Clear all matched transactions"),
            ("•", "Investigate and resolve any discrepancies before closing"),
            ("•", "Document reconciliation with statement and software screenshot"),
        ]),
        ("STEP 3 — ACCOUNTS REVIEW", [
            ("•", "Review Accounts Receivable: flag overdue invoices"),
            ("•", "Review Accounts Payable: flag upcoming bills due"),
            ("•", "Review payroll entries match payroll reports"),
            ("•", "Review loan balances match amortization schedules"),
            ("•", "Check owner's equity / draws are properly recorded"),
        ]),
        ("STEP 4 — FINANCIAL REPORT PREPARATION", [
            ("•", "Run Profit & Loss (month and YTD)"),
            ("•", "Run Balance Sheet (as of month-end)"),
            ("•", "Run Statement of Cash Flows (if applicable)"),
            ("•", "Write 3–5 bullet summary notes on key observations"),
            ("•", "Flag any items needing client decisions"),
        ]),
        ("STEP 5 — DELIVERY", [
            ("•", "Export reports as PDF"),
            ("•", "Send monthly report email using template"),
            ("•", "Log delivery date in client tracker"),
            ("•", "Note any action items for next month"),
        ]),
    ])

doc("03_BOOKKEEPING_SERVICES/Chart_of_Accounts_Guide.docx",
    "Chart of Accounts Guide",
    "Bookkeeping Business Operating System | Bookkeeping Services",
    [
        ("WHAT IS A CHART OF ACCOUNTS", [
            "The Chart of Accounts (COA) is the backbone of your client's bookkeeping system — a categorized list of every account used to track financial transactions. A well-organized COA produces clean, meaningful reports.",
        ]),
        ("STANDARD COA STRUCTURE", [
            ("•", "1000–1999: Assets (bank accounts, receivables, inventory, fixed assets)"),
            ("•", "2000–2999: Liabilities (credit cards, loans, payroll liabilities, sales tax payable)"),
            ("•", "3000–3999: Equity (owner's equity, retained earnings, owner's draws)"),
            ("•", "4000–4999: Income (sales, service revenue, interest income)"),
            ("•", "5000–5999: Cost of Goods Sold (direct materials, direct labor, subcontractors)"),
            ("•", "6000–6999: Operating Expenses (rent, utilities, marketing, software, salaries)"),
            ("•", "7000–7999: Other Income and Expenses (non-operating items)"),
        ]),
        ("INDUSTRY-SPECIFIC COA NOTES", [
            ("•", "Restaurants: Track food vs. beverage COGS separately; add 'comps and voids' expense category"),
            ("•", "Contractors / Construction: Use class/project tracking for job costing; separate labor and materials by job"),
            ("•", "E-commerce: Track COGS by product category; separate platform fees (Amazon, Shopify) from general marketing"),
            ("•", "Service Businesses: Distinguish retainer revenue vs. project revenue; track contractor vs. employee costs"),
            ("•", "Real Estate: Track rental income by property; separate depreciation, mortgage interest, repairs"),
        ]),
        ("COA BEST PRACTICES", [
            ("•", "Keep it simple: 50–100 accounts is usually sufficient; too many = hard to analyze"),
            ("•", "Use sub-accounts for detail: keep parent accounts clean"),
            ("•", "Never delete accounts with history — merge or make inactive"),
            ("•", "Match COA to how the client wants to see their business (by location, product line, department)"),
            ("•", "Align with tax return categories to make tax prep faster"),
        ]),
        ("COMMON COA MISTAKES TO FIX", [
            ("•", "Miscategorized owner draws as business expenses"),
            ("•", "Loans recorded as income"),
            ("•", "Credit card payments recorded as expenses (instead of liability payments)"),
            ("•", "Personal expenses mixed with business accounts"),
            ("•", "Opening balance equity left unresolved"),
        ]),
    ])

doc("03_BOOKKEEPING_SERVICES/Bank_Reconciliation_SOP.docx",
    "Bank Reconciliation SOP",
    "Bookkeeping Business Operating System | Bookkeeping Services",
    [
        ("PURPOSE", [
            "Bank reconciliation confirms that your accounting records match your bank statement. It catches duplicate entries, missed transactions, bank errors, and fraudulent charges. Every account should be reconciled monthly.",
        ]),
        ("STEP 1 — GATHER STATEMENTS", [
            ("•", "Download or obtain the bank/credit card statement for the month"),
            ("•", "Note the statement start and end dates and ending balance"),
            ("•", "Open the accounting software reconciliation screen for the same account"),
        ]),
        ("STEP 2 — MATCH TRANSACTIONS", [
            ("•", "Work through the statement line by line"),
            ("•", "Check off each transaction in the software that matches the statement"),
            ("•", "If a transaction appears in software but not on statement — it may be an error or timing difference"),
            ("•", "If a transaction appears on statement but not in software — add it"),
        ]),
        ("STEP 3 — RESOLVE DISCREPANCIES", [
            ("•", "Outstanding checks: issued but not yet cleared the bank — normal; list them"),
            ("•", "Outstanding deposits: posted in software but not on statement — verify bank shows it next month"),
            ("•", "Duplicate transactions: delete the duplicate, not both"),
            ("•", "Bank fees: add to books if not already recorded"),
            ("•", "Errors: if bank error, contact bank; if software error, correct the entry"),
        ]),
        ("STEP 4 — CLOSE RECONCILIATION", [
            ("•", "The 'Difference' column in your software should show $0.00"),
            ("•", "If not zero: do not force the reconciliation — investigate the difference"),
            ("•", "Take a screenshot of the completed reconciliation screen"),
            ("•", "Save the bank statement PDF to the client's folder"),
        ]),
        ("COMMON RECONCILIATION ISSUES", [
            ("•", "Transferred funds recorded twice (debit and credit not matched)"),
            ("•", "Deposits in transit from prior month not cleared"),
            ("•", "Voided checks still appearing as outstanding"),
            ("•", "Manual entries not matching actual bank amounts"),
        ]),
    ])

doc("03_BOOKKEEPING_SERVICES/Accounts_Payable_SOP.docx",
    "Accounts Payable SOP",
    "Bookkeeping Business Operating System | Bookkeeping Services",
    [
        ("PURPOSE", [
            "Accounts Payable (AP) tracks money the business owes to vendors and suppliers. Proper AP management prevents missed payments, late fees, and strained vendor relationships.",
        ]),
        ("STEP 1 — INVOICE RECEIPT AND ENTRY", [
            ("•", "Receive vendor invoices via [email / mail / portal]"),
            ("•", "Enter each invoice into accounting software: vendor, amount, date, due date, GL account"),
            ("•", "Attach the invoice PDF to the record in software"),
            ("•", "Verify invoice against purchase order or contract if applicable"),
        ]),
        ("STEP 2 — APPROVAL PROCESS", [
            ("•", "Invoices under $[X]: auto-approve and process"),
            ("•", "Invoices over $[X]: send to [owner/manager] for approval before payment"),
            ("•", "Recurring invoices: confirm amount matches contract before paying"),
        ]),
        ("STEP 3 — PAYMENT PROCESSING", [
            ("•", "Run AP aging report weekly to identify invoices due within 7 days"),
            ("•", "Process payment via [check / ACH / bill pay]"),
            ("•", "Record payment in accounting software and match to invoice"),
            ("•", "Mark invoice as paid; retain payment confirmation"),
        ]),
        ("STEP 4 — VENDOR MANAGEMENT", [
            ("•", "Maintain vendor list with payment terms, W-9 status, and 1099 requirements"),
            ("•", "Flag vendors who require 1099s (non-incorporated, paid $600+ in the year)"),
            ("•", "Track outstanding vendor credits and apply to future invoices"),
        ]),
        ("MONTH-END AP CLOSE", [
            ("•", "Review AP aging — confirm all overdue items have explanation"),
            ("•", "Reconcile AP balance on Balance Sheet to AP subledger"),
            ("•", "Review accruals for invoices received but not yet entered"),
        ]),
    ])

doc("03_BOOKKEEPING_SERVICES/Accounts_Receivable_SOP.docx",
    "Accounts Receivable SOP",
    "Bookkeeping Business Operating System | Bookkeeping Services",
    [
        ("PURPOSE", [
            "Accounts Receivable (AR) tracks money owed to the business by customers. Strong AR management improves cash flow and reduces bad debt.",
        ]),
        ("STEP 1 — INVOICE CREATION AND DELIVERY", [
            ("•", "Create invoices promptly — same day as service delivery is best"),
            ("•", "Include: client name, invoice number, itemized services, amount, due date, payment instructions"),
            ("•", "Send via [email / client portal / accounting software auto-send]"),
            ("•", "Record invoice in accounting software"),
        ]),
        ("STEP 2 — PAYMENT TRACKING", [
            ("•", "Review AR aging report weekly"),
            ("•", "Flag invoices 7+ days past due for follow-up"),
            ("•", "Apply payments received to the correct invoice immediately"),
            ("•", "Record partial payments and note remaining balance"),
        ]),
        ("STEP 3 — COLLECTIONS FOLLOW-UP SEQUENCE", [
            ("•", "Day 1 past due: Automated payment reminder (set up in accounting software)"),
            ("•", "Day 7 past due: Personal email from bookkeeper: 'Hi [Name], just a reminder that invoice #[X] for $[X] was due on [date]. Please let me know if you have questions.'"),
            ("•", "Day 14 past due: Phone call from client owner"),
            ("•", "Day 30 past due: Formal collections notice; consider pausing services"),
            ("•", "Day 60+ past due: Write off as bad debt (after client approval); consider collections agency"),
        ]),
        ("STEP 4 — CASH RECEIPT POSTING", [
            ("•", "When payment received, apply to correct open invoice"),
            ("•", "Match bank deposit to accounting software entry"),
            ("•", "If overpayment: record credit on client account or refund"),
        ]),
        ("MONTH-END AR CLOSE", [
            ("•", "Review AR aging for accuracy"),
            ("•", "Confirm AR balance on Balance Sheet matches subledger"),
            ("•", "Flag accounts over 60 days for client discussion"),
            ("•", "Review and record any bad debt write-offs"),
        ]),
    ])

print("✓ 03_BOOKKEEPING_SERVICES complete")
print("PART 1 DONE")
