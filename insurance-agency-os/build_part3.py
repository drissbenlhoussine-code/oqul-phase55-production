#!/usr/bin/env python3
"""Insurance Agency OS - Part 3: Folders 08-10"""
import os, csv
from docx import Document
from docx.shared import Pt, RGBColor

BASE = "/home/user/oqul-phase55-production/insurance-agency-os/Ultimate_Insurance_Agency_Operating_System/"

def doc(filename, title, subtitle, sections):
    d = Document()
    t = d.add_paragraph(title); t.style = d.styles['Normal']
    t.runs[0].bold = True; t.runs[0].font.size = Pt(14)
    t.runs[0].font.color.rgb = RGBColor(0x0F,0x34,0x60)
    if subtitle:
        s = d.add_paragraph(subtitle); s.style = d.styles['Normal']
        s.runs[0].font.size = Pt(10)
        s.runs[0].font.color.rgb = RGBColor(0x7F,0x8C,0x8D)
    d.add_paragraph("")
    for section in sections:
        if isinstance(section, str):
            d.add_paragraph(section); continue
        heading, items = section
        h = d.add_heading(heading, level=1)
        h.runs[0].font.color.rgb = RGBColor(0x0F,0x34,0x60)
        for item in items:
            if isinstance(item, tuple) and item[0] == '•':
                d.add_paragraph(item[1], style='List Bullet')
            else:
                d.add_paragraph(str(item))
    d.save(BASE + filename)
    print(f"  ✓ {filename}")

def write_csv(path, headers, rows):
    with open(BASE+path, "w", newline="", encoding="utf-8-sig") as f:
        w=csv.writer(f); w.writerow(headers); w.writerows(rows)
    print(f"  ✓ {path}")

# ─── 08_AGENCY_OPERATIONS ──────────────────────────────────────────────────────
p08 = "08_AGENCY_OPERATIONS/"

doc(p08+"Agency_Hiring_Guide.docx",
    "Agency Hiring Guide",
    "Insurance Agency OS | Recruiting, Hiring, and Onboarding Producers and CSRs",
    [
        ("When to Hire", [
            ("•", "Hire a CSR (Customer Service Representative) when admin tasks consume more than 25% of your week"),
            ("•", "Hire a producer when your personal pipeline exceeds your capacity to close"),
            ("•", "Consider virtual assistants before full-time hires to test the workload"),
            ("•", "Rule of thumb: Your revenue should be 3x the cost of any new hire before you hire"),
        ]),
        ("Job Descriptions", [
            "CSR Job Description Key Points:",
            ("•", "Responsibilities: Policy changes, endorsements, certificates, renewals, client service calls"),
            ("•", "Required: Active P&C license (or willingness to obtain within 90 days)"),
            ("•", "Skills: Detail-oriented, excellent phone skills, insurance software experience"),
            ("•", "Compensation: $35,000-$55,000 base + benefits; performance bonus structure"),
            "Producer Job Description Key Points:",
            ("•", "Responsibilities: Lead follow-up, needs analysis, quoting, closing, referral partner management"),
            ("•", "Required: Active P&C and/or L&H license"),
            ("•", "Skills: Sales ability, self-motivated, phone prospecting confidence"),
            ("•", "Compensation: Base draw + commission (60/40 to 80/20 base/commission split)"),
        ]),
        ("Interview Questions — Producer Candidates", [
            ("•", "Tell me about a time you overcame a significant sales objection. What did you do?"),
            ("•", "Walk me through how you would prospect for new commercial clients in this market."),
            ("•", "What's your daily prospecting routine look like?"),
            ("•", "How do you handle rejection? Give me an example of persistence paying off."),
            ("•", "What insurance lines are you most comfortable quoting, and which would you need training on?"),
            ("•", "Where do you see yourself in 3 years — managing a book of business or building a team?"),
        ]),
        ("New Hire Onboarding Checklist", [
            ("•", "Day 1: System access (AMS, email, quoting platforms), office tour, team introductions"),
            ("•", "Week 1: License verification, E&O coverage added, carrier portal access"),
            ("•", "Week 2: Shadow existing producer on 3 client meetings and 2 prospecting calls"),
            ("•", "Week 3: Role-play needs analysis and presentation — pass internal quality check"),
            ("•", "Week 4: First solo client interaction with manager review"),
            ("•", "Month 2: Assigned a small book of accounts to service and grow"),
            ("•", "Month 3: First independent new business submission"),
        ]),
    ])

doc(p08+"Agency_Training_Program.docx",
    "Agency Training Program",
    "Insurance Agency OS | 90-Day New Hire Development Plan",
    [
        ("Training Philosophy", [
            "The first 90 days with a new agent determine their trajectory in your agency. A structured training program accelerates performance, reduces mistakes, and reduces early turnover. This 90-day plan builds competency in product knowledge, sales skills, and agency systems.",
        ]),
        ("Month 1 — Foundation (Days 1-30)", [
            ("•", "Week 1: Agency systems, compliance, E&O training, carrier portals overview"),
            ("•", "Week 2: Personal lines product deep dive (home, auto, umbrella, flood)"),
            ("•", "Week 3: Sales process — needs analysis, quoting, proposal presentation"),
            ("•", "Week 4: Role-play week — complete 5 mock sales scenarios"),
            "Month 1 Milestone: Pass internal product knowledge quiz (80%+ to proceed)",
        ]),
        ("Month 2 — Application (Days 31-60)", [
            ("•", "Week 5: Live shadowing — attend 5 client appointments with senior agent"),
            ("•", "Week 6: Supervised quoting — quote 10 prospects with manager review"),
            ("•", "Week 7: First solo client calls (10 outbound calls/day, manager debrief)"),
            ("•", "Week 8: First new business submissions — 3 minimum"),
            "Month 2 Milestone: 3 submitted applications and 1 bound policy",
        ]),
        ("Month 3 — Independence (Days 61-90)", [
            ("•", "Week 9-10: Referral partner outreach — identify and contact 10 COIs"),
            ("•", "Week 11: Set personal production goals for Month 4+"),
            ("•", "Week 12: Full independence — own pipeline, own book responsibilities"),
            "Month 3 Milestone: 5 bound policies and 3 active referral partner relationships",
        ]),
        ("Ongoing Development", [
            ("•", "Weekly: 30-minute one-on-one with manager to review pipeline and skills"),
            ("•", "Monthly: Group sales training and product knowledge session"),
            ("•", "Quarterly: Performance review with goals for next quarter"),
            ("•", "Annual: Industry conference attendance and advanced certification"),
        ]),
    ])

doc(p08+"Agency_Operations_Manual.docx",
    "Agency Operations Manual",
    "Insurance Agency OS | Daily, Weekly, and Monthly Operating Procedures",
    [
        ("Daily Operating Procedures", [
            ("•", "8:00am: Review and prioritize daily task list in AMS/CRM"),
            ("•", "8:30am: Review overnight email and respond to urgent client needs first"),
            ("•", "9:00-11:00am: Peak prospecting time — outbound calls and lead follow-ups"),
            ("•", "11:00am: Review policy changes, endorsements, and service requests"),
            ("•", "1:00-3:00pm: Client appointments, presentations, and renewals"),
            ("•", "3:00pm: Process all pending policy changes and applications"),
            ("•", "4:30pm: End-of-day review: log all contacts, update CRM, set tomorrow's priorities"),
        ]),
        ("Weekly Operating Procedures", [
            ("•", "Monday: Weekly planning meeting (team), review pipeline, assign priorities"),
            ("•", "Tuesday: Focus day — prospecting calls and referral partner outreach"),
            ("•", "Wednesday: Renewal review — contact all clients with renewals in 60 days"),
            ("•", "Thursday: Client service day — annual reviews, cross-sell conversations"),
            ("•", "Friday: Administrative completion — all pending items cleared before week ends"),
        ]),
        ("Monthly Operating Procedures", [
            ("•", "Week 1: Review previous month KPIs vs. goals — identify gaps"),
            ("•", "Week 1: Carrier statement reconciliation — confirm commissions received"),
            ("•", "Week 2: Client appreciation touchpoints — thank top 10% of clients"),
            ("•", "Week 3: Referral partner check-in — lunch or coffee with top 5 COIs"),
            ("•", "Week 4: Business development — identify new carrier relationships, marketing campaigns"),
        ]),
        ("Agency Management System (AMS) Standards", [
            ("•", "Every client interaction logged within 24 hours — no exceptions"),
            ("•", "All policy changes processed and confirmed within 1 business day"),
            ("•", "Renewal diaries set 90 days in advance for every policy"),
            ("•", "Claims logged immediately upon receipt and tracked through resolution"),
            ("•", "Prospect records updated after every contact attempt"),
        ]),
    ])

print("✓ 08_AGENCY_OPERATIONS complete")

# ─── 09_POST_SALE_RETENTION ────────────────────────────────────────────────────
p09 = "09_POST_SALE_RETENTION/"

doc(p09+"Client_Retention_System.docx",
    "Client Retention System",
    "Insurance Agency OS | Keeping Clients for Life",
    [
        ("The Cost of Client Loss", [
            "Acquiring a new insurance client costs 5-7x more than retaining an existing one. A 5% improvement in retention can increase profitability by 25-95%. Building systems that proactively retain clients is the highest ROI activity in your agency.",
        ]),
        ("12-Month Client Touchpoint Calendar", [
            ("•", "Month 1 (30 days after bind): Welcome call — confirm everything is in order"),
            ("•", "Month 3: Cross-sell touchpoint — 'Let me look at your full insurance picture'"),
            ("•", "Month 6: Mid-year check-in — 'Anything changed in the past 6 months?'"),
            ("•", "Month 9: Pre-renewal outreach — 'Your renewal is coming up, let's review'"),
            ("•", "Month 12: Annual policy review — comprehensive coverage and rate review"),
            ("•", "Birthday: Send a card or text on the client's birthday"),
            ("•", "Policy anniversary: Note the anniversary and thank them for their loyalty"),
        ]),
        ("VIP Client Program", [
            "Identify your top 20% of clients (by premium and referral production) and give them VIP treatment:",
            ("•", "Designated direct line to a senior agent — no hold times"),
            ("•", "Annual in-person or video review (not just a phone call)"),
            ("•", "Priority claims handling — agent personally advocates from day one"),
            ("•", "Annual VIP client event invitation — dinner, golf, entertainment"),
            ("•", "Holiday gift — thoughtful, personalized gift each December"),
            ("•", "Early access to new products and programs"),
        ]),
        ("Cancellation Prevention Protocol", [
            "When a client calls to cancel:",
            '"[Name], I\'m sorry to hear that — before I process this, would you mind telling me what\'s driving the decision? I want to make sure I have every opportunity to keep your business and that we\'ve given you our best."',
            "Listen — common reasons and responses:",
            ("•", "Price: 'Let me run a market comparison right now — I may be able to match or beat that quote'"),
            ("•", "Service issue: 'I appreciate you telling me. Tell me what happened and I want to make it right'"),
            ("•", "Moving: 'Can I get you coverage in your new state? We may have carrier appointments there'"),
            ("•", "Switching to bundle elsewhere: 'Before you go, let me see if I can match the bundle pricing — it takes 10 minutes'"),
        ]),
    ])

doc(p09+"Annual_Review_Scripts.docx",
    "Annual Review Scripts",
    "Insurance Agency OS | Comprehensive Annual Policy Review Framework",
    [
        ("Opening an Annual Review Conversation", [
            '"Hi [Name], this is [Agent] from [Agency]. It\'s been about a year since we got your coverage set up, and I make it a point to connect with every client annually to make sure your policies still make sense for your life and that you\'re getting the best value. Do you have about 15 minutes?"',
        ]),
        ("Annual Review Discovery Questions", [
            "Life Changes:",
            ("•", "'Has anything significant changed in your life this past year — job changes, moves, family changes?'"),
            ("•", "'Did you add any vehicles, drivers, or properties?'"),
            ("•", "'Any major renovations or improvements to your home?'"),
            "Coverage Assessment:",
            ("•", "'Have you had any claims or incidents this year, even if you didn't file them?'"),
            ("•", "'Do you feel like you understand what's covered and what's not on your current policies?'"),
            ("•", "'Is there any area of your financial life that feels underprotected?'"),
            "Rate Review:",
            ("•", "'I've been keeping an eye on your rates and wanted to let you know [rates have changed / I've shopped the market for you].'"),
        ]),
        ("Cross-Sell During Annual Review", [
            "At some point in every annual review, introduce this question:",
            '"One thing I like to do in these conversations is make sure we\'re looking at your complete insurance picture. Besides what we handle here, do you have [life insurance / umbrella / long-term disability / flood] coverage in place?"',
            "For each 'no' or 'I\'m not sure':",
            '"That\'s actually something I\'d love to talk through with you, because [specific risk/gap]. Do you mind if I get you a quick quote on that before we finish up today?"',
        ]),
        ("Renewal Confirmation and Close", [
            '"Based on our conversation, it sounds like your coverage is [in good shape / needs some adjustments]. Here\'s what I\'d recommend we do: [Summary of changes or continuation]. Does that all make sense to you?"',
            '"Your renewal is coming up on [Date]. I\'ll have [an updated proposal / a confirmation of your renewal] to you by [Date]. Sound good?"',
        ]),
    ])

doc(p09+"Client_Appreciation_Program.docx",
    "Client Appreciation Program",
    "Insurance Agency OS | Events, Gifts, and Recognition to Build Loyalty",
    [
        ("Why Client Appreciation Drives Revenue", [
            "Clients who feel genuinely appreciated refer more, stay longer, and buy more. An intentional appreciation program is not a cost — it's an investment in your highest-revenue source: word-of-mouth referrals from loyal clients.",
        ]),
        ("Annual Client Event Ideas", [
            ("•", "Spring: Easter celebration + insurance Q&A ('Ask an Agent' format)"),
            ("•", "Summer: Customer Appreciation BBQ or outdoor movie night"),
            ("•", "Fall: Pumpkin patch outing for families, or wine and cheese evening"),
            ("•", "Holiday: Annual holiday party (December) — dinner, giveaways, best-dressed contest"),
            ("•", "Sports: Invite top clients to a local sporting event (baseball, soccer, hockey)"),
        ]),
        ("Milestone Recognition Program", [
            ("•", "5-Year Client: Personalized 'Thank You for 5 Years' card + gift card"),
            ("•", "10-Year Client: Premium thank-you gift + featured in newsletter"),
            ("•", "25+ Policies: Recognized as agency's top client, invited to exclusive events"),
            ("•", "Major Life Event (marriage, baby, retirement): Personal note + small gift"),
        ]),
        ("Low-Cost High-Impact Appreciation Touches", [
            ("•", "Handwritten birthday cards — stand out in a world of automated emails"),
            ("•", "Policy anniversary note: 'It's been [X] years — thanks for your trust!'"),
            ("•", "Share wins: 'I got you a rate reduction at renewal — wanted to share the good news!'"),
            ("•", "Educational emails: 'I thought of you when I saw this article on [relevant topic]'"),
            ("•", "Holiday cards: A physical card in December creates lasting impressions"),
        ]),
    ])

print("✓ 09_POST_SALE_RETENTION complete")

# ─── 10_NOTION_WORKSPACE ──────────────────────────────────────────────────────
p10 = "10_NOTION_WORKSPACE/"

write_csv(p10+"Lead_Database.csv",
    ["Lead ID","Name","Phone","Email","Source","Line of Business","Status","Est. Premium","Last Contact","Next Action","Notes","Agent"],
    [
        ["L001","James Parker","(555) 201-3344","jparker@email.com","Zillow Ad","Home+Auto","Hot","$2,400","2024-01-10","Send quote","Pre-qualified, good credit","Agent Smith"],
        ["L002","Susan Hill","(555) 312-4455","shill@email.com","Referral","Auto","Warm","$1,100","2024-01-11","Follow-up call","Single mom, price-sensitive","Agent Smith"],
        ["L003","Robert Chen","(555) 423-5566","rchen@email.com","LinkedIn","Commercial BOP","Hot","$8,500","2024-01-12","Appointment set","Restaurant owner, 12 employees","Agent Jones"],
        ["L004","Maria Gonzalez","(555) 534-6677","mgonzalez@email.com","Facebook","Home+Auto+Life","Warm","$4,200","2024-01-13","Needs analysis","Young family, life insurance interest","Agent Smith"],
        ["L005","David Kim","(555) 645-7788","dkim@email.com","Cold Call","GL Only","Cold","$3,200","2024-01-14","Email follow-up","Skeptical about switching","Agent Jones"],
        ["L006","Jennifer White","(555) 756-8899","jwhite@email.com","SOI","Full Package","Hot","$5,100","2024-01-15","Quote presentation","Neighbor of existing client","Agent Smith"],
        ["L007","Carlos Rivera","(555) 867-9900","crivera@email.com","Chamber","Commercial","Warm","$6,800","2024-01-16","Send questionnaire","Construction company owner","Agent Jones"],
        ["L008","Amy Johnson","(555) 978-0011","ajohnson@email.com","Website","Auto","Cold","$900","2024-01-17","Send email","Just browsing rates","Agent Smith"],
    ])

write_csv(p10+"Client_Database.csv",
    ["Client ID","Name","Phone","Email","Policy Types","Annual Premium","Carrier(s)","Next Renewal","Tier","Last Contact","Cross-Sell Opportunity","Notes"],
    [
        ["C001","Thomas Brown","(555) 111-2222","tbrown@email.com","Home+Auto","$2,850","Travelers","2024-06-15","A","2024-01-10","Umbrella","Loyal 7-year client, strong referrer"],
        ["C002","Patricia Davis","(555) 222-3333","pdavis@email.com","Home+Auto+Umbrella","$4,100","Nationwide","2024-07-01","VIP","2024-01-05","Life Insurance","Top referrer — sent 5 clients last year"],
        ["C003","Michael Lee","(555) 333-4444","mlee@email.com","BOP+WC+Auto","$11,200","Hartford","2024-08-15","VIP","2023-08-10","Cyber Liability","Commercial client growing rapidly"],
        ["C004","Sandra Wilson","(555) 444-5555","swilson@email.com","Home+Auto+Life","$5,400","Progressive","2024-05-30","A","2023-12-15","Disability","Recently married, life added"],
        ["C005","Kevin Thompson","(555) 555-6666","kthompson@email.com","Auto Only","$1,100","GEICO","2024-09-01","B","2024-01-08","Home+Bundle","Renter now — considering buying"],
        ["C006","Linda Martinez","(555) 666-7777","lmartinez@email.com","Home+Auto+Flood","$3,900","Allstate","2024-04-15","A","2023-10-15","Jewelry Rider","Coastal property, flood added last year"],
    ])

write_csv(p10+"Claims_Tracker.csv",
    ["Claim #","Client Name","Policy Type","Date Filed","Carrier","Loss Type","Status","Settlement","Days Open","Adjuster","Last Update","Agent Notes"],
    [
        ["CLM-001","Thomas Brown","Auto","2024-01-05","Travelers","Rear-end collision","Open","TBD","15","Mike Torres","Adjuster inspected vehicle 1/12","Awaiting repair estimate"],
        ["CLM-002","Patricia Davis","Home","2023-12-20","Nationwide","Wind damage","Settled","$8,500","40","Sarah Chen","Check issued 1/29","Client satisfied with outcome"],
        ["CLM-003","Linda Martinez","Home+Flood","2024-01-10","Allstate","Water damage","Open","TBD","10","James Wu","Inspector scheduled 1/18","Flood adjuster separate from homeowner adjuster"],
        ["CLM-004","Michael Lee","WC","2023-11-15","Hartford","Employee injury","Open","$24,000 reserve","65","Tom Hill","IME scheduled 2/1","Employee back to work light duty"],
        ["CLM-005","Sandra Wilson","Auto","2023-12-28","Progressive","Parking lot fender bender","Settled","$1,200","22","Lena Park","Closed 1/19","Minor claim, no rate impact expected"],
    ])

write_csv(p10+"Referral_Partner_Database.csv",
    ["Partner ID","Name","Company","Type","Phone","Email","Sent","Received","Relationship","Last Meeting","Notes"],
    [
        ["RP001","David Reyes","First Alliance Mortgage","Mortgage Broker","(555) 100-2200","dreyes@fam.com","12","8","Strong","2024-01-08","Monthly lunch, very reliable referrer"],
        ["RP002","Christine Park","Park Realty Group","Real Estate Agent","(555) 200-3300","cpark@parkrealty.com","18","6","Strong","2024-01-05","Top partner, 6 closed referrals last year"],
        ["RP003","Brian Foster","Foster Financial","Financial Advisor","(555) 300-4400","bfoster@ff.com","7","4","Developing","2023-12-20","Quarterly meetings, growing relationship"],
        ["RP004","Angela Moore","Moore CPA Group","CPA / Tax","(555) 400-5500","amoore@moorecpa.com","5","2","Developing","2023-11-15","Good for commercial referrals"],
        ["RP005","Marco Silva","Silva Law Firm","Attorney","(555) 500-6600","msilva@silvalaw.com","4","1","New","2024-01-10","Met at chamber event, promising"],
        ["RP006","Kelly Zhang","Zhang Auto Group","Car Dealership","(555) 600-7700","kzhang@zhangauto.com","22","11","Strong","2024-01-12","Best volume source for auto leads"],
    ])

write_csv(p10+"Policy_Renewal_Queue.csv",
    ["Policy #","Client Name","Line","Carrier","Renewal Date","Current Premium","Status","Action","Priority","Agent"],
    [
        ["POL-001","Linda Martinez","Home+Auto+Flood","Allstate","2024-04-15","$3,900","Shop Market","Call by 3/15","HIGH","Agent Smith"],
        ["POL-002","Sandra Wilson","Home+Auto+Life","Progressive","2024-05-30","$5,400","Review","Annual review scheduled 5/1","MEDIUM","Agent Smith"],
        ["POL-003","Thomas Brown","Home+Auto","Travelers","2024-06-15","$2,850","Auto-Renew","Confirm + cross-sell umbrella","MEDIUM","Agent Smith"],
        ["POL-004","Patricia Davis","Home+Auto+Umbrella","Nationwide","2024-07-01","$4,100","Review","VIP review — in-person meeting","HIGH","Agent Smith"],
        ["POL-005","Kevin Thompson","Auto","GEICO","2024-09-01","$1,100","Cross-Sell","Upsell opportunity — buying a home","MEDIUM","Agent Jones"],
        ["POL-006","Michael Lee","BOP+WC+Auto","Hartford","2024-08-15","$11,200","Loss Runs","Request 5-yr loss runs by 7/1","HIGH","Agent Jones"],
    ])

write_csv(p10+"Producer_Activity_Log.csv",
    ["Date","Agent","Activity Type","Contact","Result","Follow-Up Date","Notes"],
    [
        ["2024-01-10","Agent Smith","Outbound Call","James Parker","Quote requested","2024-01-12","Ready to move forward, just needs quote"],
        ["2024-01-10","Agent Smith","Annual Review","Thomas Brown","Cross-sell opportunity","2024-02-15","Interested in umbrella — send info"],
        ["2024-01-11","Agent Jones","LinkedIn Outreach","Robert Chen","Appointment set","2024-01-18","Commercial BOP for restaurant, 12 employees"],
        ["2024-01-12","Agent Smith","Open House Event","Jennifer White","Referral from neighbor","2024-01-15","Needs full package — home, auto, umbrella"],
        ["2024-01-13","Agent Jones","Cold Call","David Kim","Left voicemail","2024-01-17","Second follow-up attempt — call again"],
        ["2024-01-14","Agent Smith","Referral Partner Lunch","Christine Park","3 referrals promised","2024-02-01","Monthly lunch, great relationship"],
        ["2024-01-15","Agent Smith","Client Service","Patricia Davis","Endorsement processed","None","Added jewelry rider, confirmed via email"],
        ["2024-01-16","Agent Jones","Commercial Appointment","Carlos Rivera","Needs analysis complete","2024-01-22","Submitting BOP and commercial auto to 3 carriers"],
    ])

write_csv(p10+"Commission_Tracker.csv",
    ["Month","Carrier","Policy Type","Policies","Premium","Commission %","Commission Earned","Contingency","Total Revenue"],
    [
        ["January 2024","Travelers","Home+Auto","22","$52,800","12%","$6,336","$0","$6,336"],
        ["January 2024","Nationwide","Home+Auto+Umbrella","15","$35,200","11%","$3,872","$0","$3,872"],
        ["January 2024","Hartford","Commercial","8","$89,500","10%","$8,950","$500","$9,450"],
        ["January 2024","Progressive","Auto","30","$28,100","13%","$3,653","$0","$3,653"],
        ["February 2024","Travelers","Home+Auto","25","$61,000","12%","$7,320","$0","$7,320"],
        ["February 2024","Hartford","Commercial","11","$112,000","10%","$11,200","$700","$11,900"],
        ["February 2024","Nationwide","Various","18","$41,500","11%","$4,565","$0","$4,565"],
        ["February 2024","Progressive","Auto","28","$25,800","13%","$3,354","$0","$3,354"],
    ])

# Markdown Guide
md = """# Insurance Agency OS — Notion Workspace Setup Guide

## Overview
This guide walks you through importing your 7 Insurance Agency OS databases into Notion, creating your complete digital agency management system in under 30 minutes.

## Files Included
| File | Notion Database | Purpose |
|------|----------------|---------|
| Lead_Database.csv | Lead Database | Track all incoming prospects |
| Client_Database.csv | Client Database | Manage your book of business |
| Claims_Tracker.csv | Claims Tracker | Monitor open and closed claims |
| Referral_Partner_Database.csv | Partner Network | Manage COI relationships |
| Policy_Renewal_Queue.csv | Renewal Queue | Never miss a renewal |
| Producer_Activity_Log.csv | Activity Log | Track daily sales activities |
| Commission_Tracker.csv | Commission Tracker | Monitor revenue by carrier |

## Import Instructions

### Step 1 — Create Your Agency Hub
1. Open Notion and click **+ New Page**
2. Title it: "Insurance Agency OS"
3. Add icon: 🛡️ and a professional cover image

### Step 2 — Import Each Database
For each CSV file:
1. Click **+ New Page** inside your Agency OS page
2. Select **Import** → **CSV**
3. Upload the CSV file → Click **Import**
4. Notion creates a table database automatically

### Step 3 — Recommended Views Per Database

**Lead Database:**
- Table View (default): All leads
- Board View: Group by Status (Hot/Warm/Cold)
- Filter View: "Hot Leads" — Status = Hot, sorted by Last Contact
- Filter View: "Follow Up Today" — Next Action date = Today

**Policy Renewal Queue:**
- Table View: All renewals
- Board View: Group by Priority (HIGH/MEDIUM/LOW)
- Calendar View: Group by Renewal Date
- Filter View: "Urgent" — Priority = HIGH

**Claims Tracker:**
- Table View: All claims
- Board View: Group by Status (Open/Settled)
- Filter View: "Open Claims" — Status = Open, sorted by Days Open

### Step 4 — Link Your Databases
1. In **Claims Tracker**: Add Relation → Client Database
2. In **Policy Renewal Queue**: Add Relation → Client Database
3. In **Producer Activity Log**: Add Relation → Lead Database

### Step 5 — Daily Dashboard Page
Create a "Daily Dashboard" page with linked views:
- Hot Leads (filtered from Lead Database)
- Today's Follow-Ups (filtered from Activity Log)
- Urgent Renewals (filtered from Renewal Queue)
- Open Claims requiring action

## Pro Tips
- Log every client contact in Producer Activity Log same day
- Update Lead Status immediately after each conversation
- Set Renewal Queue reminders 90 days in advance
- Review Commission Tracker monthly to verify carrier payments

---
*Insurance Agency Operating System | Professional Edition*
*Original Price: €149 | Your Price: €39*
"""
with open(BASE+p10+"Notion_Setup_Guide.md","w",encoding="utf-8") as f:
    f.write(md)
print(f"  ✓ {p10}Notion_Setup_Guide.md")
print("✓ 10_NOTION_WORKSPACE complete")

print("\nPART 3 DONE")
